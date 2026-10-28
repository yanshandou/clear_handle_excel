import importlib
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl
from openpyxl.styles import Color, Font, PatternFill
from types import SimpleNamespace

from core import fix_images, mark_finder, reporter, storage, workflow

MAGENTA = [{"fill": "#C0007A", "font": "#FFFFFF"}, {"fill": "#000000", "font": "#FFFF00"}]


def paint(cell, fill, font=None):
    cell.fill = PatternFill("solid", fgColor="FF" + fill)
    if font:
        cell.font = Font(color="FF" + font)


class MarkFinderTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        for obj, name, value in ((storage, "DATABASE_PATH", self.root / "state.sqlite3"),
                                 (reporter, "RECORD_PATH", self.root / "record.md"),
                                 (fix_images, "IMAGES_DIR", self.root / "images")):
            p = patch.object(obj, name, value)
            p.start()
            self.addCleanup(p.stop)
        self.wf = importlib.reload(workflow)
        mark_finder.wf = self.wf
        self.folder = self.root / "成品"
        self.folder.mkdir()
        self.path = self.folder / "book.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "明細"
        paint(sheet["A5"], "C0007A", "FFFFFF")       # 底色字色都对上
        paint(sheet["B7"], "C0007A", "123456")       # 只有底色对上（一格几种字色时旧版不改字色）
        paint(sheet["C9"], "C0007B", "FFFFFF")       # 差一点也不算
        paint(sheet["D3"], "FFFF00")                 # 别的颜色
        sheet.merge_cells("E2:G2")
        for col in "EFG":
            paint(sheet[f"{col}2"], "000000", "FFFF00")  # 合并区域整块算一格
        book.save(self.path)
        (self.folder / "old.xls").write_bytes(b"x")

    def test_exact_marker_colors_are_listed_merged_area_once(self):
        hits = mark_finder.scan_file(self.path, mark_finder._styles(MAGENTA), set())
        found = {h["addr"]: h["match"] for h in hits}
        self.assertEqual(found, {"A5": "both", "B7": "fill", "E2:G2": "both"})
        self.assertEqual(next(h for h in hits if h["addr"] == "B7")["font"], "#123456")

    def test_cells_the_program_still_has_records_for_are_flagged(self):
        self.wf.STATE.setdefault("batch_review", {})["highlights"] = {
            "x": {"path": str(self.path), "sheet": "明細", "addr": "$A$5", "cells": []},
            "y": {"path": str(self.path), "sheet": "明細", "addr": "$B$7", "tag": "CHE_MARK_y"}}  # 新方式的不算
        hits = mark_finder.scan_file(self.path, mark_finder._styles(MAGENTA), mark_finder._recorded(str(self.path)))
        self.assertEqual({h["addr"] for h in hits if h["recorded"]}, {"A5"})

    def test_history_with_same_fill_and_different_fonts_matches_every_pair(self):
        colors = MAGENTA + [{"fill": "#C0007A", "font": "#123456"}]
        hits = mark_finder.scan_file(self.path, mark_finder._styles(colors), set())
        self.assertEqual({h["addr"]: h["match"] for h in hits},
                         {"A5": "both", "B7": "both", "E2:G2": "both"})

    def test_background_run_reports_each_file_and_unreadable_formats(self):
        with patch.object(self.wf, "log_msg") as log, \
                patch("core.excel_ops.start_background_excel", side_effect=RuntimeError("没有 Excel")):
            mark_finder.start(str(self.folder), MAGENTA)
            for _ in range(100):
                if not mark_finder.status()["running"]:
                    break
                time.sleep(0.05)
        result = mark_finder.status()["result"]
        self.assertEqual(result["total"], 2)
        files = {f["name"]: f for f in result["files"]}
        self.assertEqual(len(files["book.xlsx"]["hits"]), 3)
        self.assertIn("用 Excel 读取失败", files["old.xls"]["error"])
        self.assertTrue(any("3 格底色是标记色" in c.args[0] for c in log.call_args_list))

    def old_copy(self, edit=None):
        """旧版文件：同样的内容，原来的颜色"""
        old_root = self.root / "原始备份"
        old_root.mkdir(exist_ok=True)
        book = openpyxl.load_workbook(self.path)
        sheet = book["明細"]
        for addr in ("A5", "B7", "E2", "F2", "G2"):
            sheet[addr].fill = PatternFill(fill_type=None)
            sheet[addr].font = Font(color="FF000000")
        paint(sheet["B7"], "DDEEFF", "123456")
        if edit:
            edit(sheet)
        book.save(old_root / "book.xlsx")
        return old_root

    def plan(self, old_root):
        hits = mark_finder.scan_file(self.path, mark_finder._styles(MAGENTA), set())
        return {c["addr"]: c for c in mark_finder.plan_file(self.path, old_root / "book.xlsx", hits)}

    def test_plan_restores_cells_whose_content_matches_the_old_version(self):
        book = openpyxl.load_workbook(self.path)
        book["明細"]["A5"] = "合計"
        book["明細"]["B7"] = 12
        book.save(self.path)
        plan = self.plan(self.old_copy())
        self.assertEqual((plan["A5"]["action"], plan["A5"]["old_fill"], plan["A5"]["old_font"]), ("restore", "none", "#000000"))
        self.assertEqual((plan["B7"]["action"], plan["B7"]["old_fill"]), ("restore", "#DDEEFF"))
        self.assertEqual(plan["E2:G2"]["action"], "restore")

    def test_plan_skips_moved_or_changed_cells_and_original_marker_colors(self):
        book = openpyxl.load_workbook(self.path)
        book["明細"]["A5"] = "合計"
        book["明細"]["B7"] = "x"
        book.save(self.path)

        def edit(sheet):
            sheet["A5"] = "小計"                 # 内容不同（多半是插了行）
            sheet["B7"] = "x"
            sheet["B8"] = "多了一格"             # 不影响 B7：B7 有内容，只比这一格
            sheet["H2"] = "这一行多了内容"       # E2 是空格子：这一行和旧版对不上
            paint(sheet["B7"], "C0007A", "123456")  # 旧版本来就是这个颜色
        plan = self.plan(self.old_copy(edit))
        self.assertIn("内容不一样", plan["A5"]["reason"])
        self.assertIn("本来就是这个颜色", plan["B7"]["reason"])
        self.assertIn("这一行和旧版对不上", plan["E2:G2"]["reason"])
        self.assertTrue(all(c["action"] == "skip" for c in plan.values()))

    def test_renamed_sheets_are_matched_by_guessed_name_code_name_then_content(self):
        book = openpyxl.load_workbook(self.path)
        book["明細"]["A5"] = "合計"
        other = book.create_sheet("レイアウト")
        other.sheet_properties.codeName = "Sheet9"
        other["A1"], other["A2"] = "訪問看護計画書", "氏名"
        paint(other["A2"], "C0007A", "FFFFFF")
        bill = book.create_sheet("請求書")
        bill["C3"] = "金額"
        paint(bill["C3"], "C0007A", "FFFFFF")
        book["明細"].sheet_properties.codeName = "Sheet1"
        book.save(self.path)

        def edit(sheet):
            sheet.title = "明細新"                 # 精简 Sheet 去掉了「新」，代码名没变
            sheet.sheet_properties.codeName = "Sheet1"
            lay = sheet.parent["レイアウト"]
            lay.title = "レイアウト(訪問看護計画書)"
            lay.sheet_properties.codeName = "Sheet3"   # 代码名也对不上，只能按内容
            lay["A2"].fill = PatternFill(fill_type=None)
            sheet.parent.create_sheet("全然違う")["A1"] = "別の表"
            bill = sheet.parent["請求書"]
            bill.title = "新請求書"                  # 「新」字被去掉了：加回去就能对上
            bill["C3"].fill = PatternFill(fill_type=None)
        plan = self.plan(self.old_copy(edit))
        self.assertEqual((plan["A5"]["action"], plan["A5"]["old_sheet"]), ("restore", "明細新"))
        self.assertIn("按代码名对上", plan["A5"]["reason"])
        self.assertEqual((plan["A2"]["action"], plan["A2"]["old_sheet"]), ("restore", "レイアウト(訪問看護計画書)"))
        self.assertIn("按内容对上", plan["A2"]["reason"])
        self.assertEqual((plan["C3"]["action"], plan["C3"]["old_sheet"]), ("restore", "新請求書"))
        self.assertIn("去掉「新」", plan["C3"]["reason"])

    def test_sheet_without_any_match_lists_old_sheet_names(self):
        def edit(sheet):
            sheet.title = "別の名前"
            sheet.sheet_properties.codeName = "Sheet7"
            for addr in ("A1", "A5", "B7"):
                sheet[addr] = "まったく違う内容"
        plan = self.plan(self.old_copy(edit))
        self.assertEqual(plan["A5"]["action"], "skip")
        self.assertIn("別の名前", plan["A5"]["reason"])

    def test_files_are_paired_by_path_then_by_unique_name(self):
        old = self.root / "old"
        paths = [old / "a.xlsx", old / "x" / "b.xlsx", old / "y" / "c.xlsx", old / "z" / "c.xlsx"]
        pairs = mark_finder._pair_files(old, paths, ["a.xlsx", "sub/b.xlsx", "c.xlsx", "d.xlsx"])
        self.assertEqual(pairs, {"a.xlsx": paths[0], "sub/b.xlsx": paths[1], "c.xlsx": None, "d.xlsx": None})

    def test_apply_copies_old_colors_only_when_still_marked_and_same_content(self):
        marker = mark_finder._excel_color("#C0007A")

        def cell(formula, pattern, color, font, tint=0):
            return Mock(Formula=formula, Interior=Mock(Pattern=pattern, Color=color, TintAndShade=tint), Font=Mock(Color=font))
        cur = {"A5": cell("合計", 1, marker, 0xFFFFFF), "B7": cell("x", 1, 0x123456, 0), "C9": cell("新", 1, marker, 0)}
        old = {"A5": cell("合計", -4142, 0xFFFFFF, 0), "B7": cell("x", -4142, 0xFFFFFF, 0),
               "C9": cell("旧", -4142, 0xFFFFFF, 0)}
        for c in cur.values():
            c.Cells.side_effect = lambda r, col, c=c: c
        old_sheet, sheet = Mock(Name="明細"), Mock()
        old_sheet.Range.side_effect = lambda a: old[a]
        sheet.Range.side_effect = lambda a: cur[a]
        old_book, book = Mock(Worksheets=[old_sheet]), Mock(ReadOnly=False)
        book.Worksheets.return_value = sheet
        excel = Mock()
        opened = []

        def open_book(path, **_):
            # 旧版和新版同名：Excel 里已经开着同名文件时 Open 返回空对象，所以旧版要先读完关掉
            if opened and not opened[-1].Close.called:
                return None
            opened.append(old_book if not opened else book)
            return opened[-1]
        excel.Workbooks.Open.side_effect = open_book
        old_sheet.Name = "明細新"  # 旧版表名不同：按清单里对上的旧表名去读
        cells = [{"sheet": "明細", "old_sheet": "明細新", "addr": a, "fill": "#C0007A", "font": "#FFFFFF", "action": "restore"} for a in cur]
        with patch("core.excel_ops.normalize_workbook_view"):
            done = {c["addr"]: c for c in mark_finder.apply_file(excel, str(self.path), str(self.path), cells)}
        self.assertEqual(done["A5"]["status"], "ok")
        self.assertEqual((cur["A5"].Interior.Pattern, cur["A5"].Font.Color), (-4142, 0))
        self.assertIn("被人改过", done["B7"]["reason"])
        self.assertIn("内容和旧版不一样", done["C9"]["reason"])
        self.assertEqual(cur["C9"].Interior.Color, marker)
        book.Save.assert_called_once()

    def test_bad_colors_are_refused(self):
        with self.assertRaises(ValueError):
            mark_finder._styles([{"fill": "red", "font": "#FFFFFF"}])

    def test_color_file_listing_includes_binary_and_templates_and_ignores_lock_files(self):
        for name in ("binary.xlsb", "template.xltx", "macro_template.xltm", "~$binary.xlsb", "notes.txt"):
            (self.folder / name).touch()
        _, paths = mark_finder.list_workbooks(str(self.folder))
        self.assertEqual({path.name for path in paths}, {"book.xlsx", "old.xls", "binary.xlsb", "template.xltx", "macro_template.xltm"})
        self.assertTrue(mark_finder._is_template("file.xltm"))
        self.assertFalse(mark_finder._is_template("file.xlsx"))

    def test_templates_are_listed_for_manual_repair_and_never_opened_for_writing(self):
        template, backup = self.folder / "marked.xltx", self.root / "original.xltx"
        book = openpyxl.Workbook()
        book.template = True
        book.save(backup)
        book.active.sheet_properties.tabColor = "FFC000"
        book.save(template)
        book.close()
        hits = mark_finder.scan_file(template, mark_finder._styles(MAGENTA), set())
        plan = mark_finder.plan_file(template, backup, hits)
        self.assertEqual(plan[0]["action"], "skip")
        self.assertIn("模板", plan[0]["reason"])
        excel = Mock()
        with self.assertRaisesRegex(ValueError, "模板"):
            mark_finder.apply_file(excel, str(template), str(backup), plan)
        excel.Workbooks.Open.assert_not_called()

    def test_tab_only_markers_restore_none_theme_and_renamed_sheet_and_preserve_original_orange(self):
        path, old_path = self.folder / "tabs.xlsx", self.root / "tabs_old.xlsx"
        book = openpyxl.Workbook()
        book.active.title = "NoColor"
        book.create_sheet("Theme").sheet_properties.tabColor = Color(theme=5, tint=0.4)
        book.create_sheet("OriginallyOrange").sheet_properties.tabColor = "FFC000"
        book.create_sheet("Manual").sheet_properties.tabColor = "112233"
        book.create_sheet("Renamed").sheet_properties.tabColor = "123456"
        for sheet in book.worksheets:
            sheet["A1"] = sheet.title
            sheet.sheet_properties.codeName = f"Sheet{len(sheet.title)}{sheet.title}"
        book.save(old_path)
        for sheet in book.worksheets:
            sheet.sheet_properties.tabColor = "FF0000" if sheet.title == "Manual" else "FFC000"
        book["Renamed"].title = "新Renamed"
        book.create_sheet("Unpaired").sheet_properties.tabColor = "FFC000"
        book.save(path)
        book.close()
        hits = mark_finder.scan_file(path, mark_finder._styles(MAGENTA), set())
        self.assertTrue(all(hit["kind"] == "tab" for hit in hits))
        self.assertNotIn("Manual", [hit["sheet"] for hit in hits])
        plan = {item["sheet"]:item for item in mark_finder.plan_file(path, old_path, hits)}
        self.assertEqual((plan["NoColor"]["action"], plan["NoColor"]["old_fill"]), ("restore", "none"))
        self.assertEqual((plan["Theme"]["action"], plan["Theme"]["old_fill"]), ("restore", "other"))
        self.assertEqual(plan["OriginallyOrange"]["action"], "skip")
        self.assertEqual(plan["新Renamed"]["old_sheet"], "Renamed")
        self.assertEqual(plan["Unpaired"]["action"], "skip")

    def tab_apply(self, original, current_color="#FFC000", include_cell=False, *,
                  versions=None, on_open=None, before_save=None):
        old_sheet = Mock(Name="Main", Tab=original)
        tab = SimpleNamespace(ColorIndex=6, Color=mark_finder._excel_color(current_color), ThemeColor=None, TintAndShade=0.0)
        current = Mock(Tab=tab)
        self.last_tab_sheet = current
        old_cell = Mock(Formula="same", Interior=Mock(Pattern=-4142, Color=0, TintAndShade=0, PatternColor=0), Font=Mock(Color=0))
        cell = Mock(Formula="same", Interior=Mock(Pattern=1, Color=mark_finder._excel_color("#C0007A")), Font=Mock())
        cell.Cells.return_value = cell
        old_sheet.Range.return_value = old_cell
        current.Range.return_value = cell
        old, book = Mock(Worksheets=[old_sheet]), Mock(ReadOnly=False)
        self.last_tab_book = book
        book.Worksheets.return_value = current
        excel = Mock()

        def open_book(path, **options):
            if options.get("ReadOnly"):
                return old
            self.assertEqual(Path(path), self.path.resolve())
            old.Close.assert_called_once_with(SaveChanges=False)
            if on_open is not None:
                on_open()
            return book

        excel.Workbooks.Open.side_effect = open_book
        items = [mark_finder._tab_hit("Main")]
        if include_cell:
            items.append({"sheet":"Main", "addr":"A1", "fill":"#C0007A", "font":"#FFFFFF"})
        with patch("core.excel_ops.normalize_workbook_view", side_effect=before_save) as normalize:
            results = mark_finder.apply_file(excel, str(self.path), str(self.root / "old.xlsx"), items,
                                            versions=versions)
        old.Close.assert_called_once_with(SaveChanges=False)
        self.assertTrue(excel.Workbooks.Open.call_args_list[0].kwargs["ReadOnly"])
        return results, book, tab, normalize

    def test_tab_restore_preserves_theme_and_clears_only_program_orange_without_changing_view(self):
        for original in (SimpleNamespace(ColorIndex=-4142),
                         SimpleNamespace(ColorIndex=5, Color=0x123456, ThemeColor=5, TintAndShade=0.4)):
            results, book, tab, normalize = self.tab_apply(original)
            self.assertEqual(results[0]["status"], "ok")
            if original.ColorIndex == -4142:
                self.assertEqual(tab.ColorIndex, -4142)
            else:
                self.assertEqual((tab.ThemeColor, tab.TintAndShade), (5, 0.4))
            book.Save.assert_called_once()
            normalize.assert_not_called()
        results, book, tab, _ = self.tab_apply(SimpleNamespace(ColorIndex=-4142), current_color="#FF0000")
        self.assertIn("保留当前颜色", results[0]["reason"])
        self.assertEqual(tab.Color, mark_finder._excel_color("#FF0000"))
        book.Save.assert_not_called()

    def test_tab_and_cell_at_a1_keep_separate_backup_data_and_both_restore(self):
        results, book, tab, normalize = self.tab_apply(SimpleNamespace(ColorIndex=-4142), include_cell=True)
        self.assertEqual([item["status"] for item in results], ["ok", "ok"])
        self.assertEqual(tab.ColorIndex, -4142)
        book.Save.assert_called_once()
        normalize.assert_called_once()

    def test_tab_write_failure_never_saves_partial_changes(self):
        with patch.object(mark_finder.batch, "_restore_tab", side_effect=RuntimeError("工作簿受保护")):
            with self.assertRaisesRegex(RuntimeError, "受保护"):
                self.tab_apply(SimpleNamespace(ColorIndex=-4142))
        self.last_tab_book.Save.assert_not_called()
        self.last_tab_book.Close.assert_called_once_with(SaveChanges=False)

    def tab_apply_versions(self):
        from core.link_repair import fingerprint

        backup = self.root / "old.xlsx"
        backup.write_bytes(b"original backup")
        return fingerprint(self.path), fingerprint(backup)

    def test_version_changes_during_current_open_abort_before_any_sheet_write(self):
        for changed_path in (self.path, self.root / "old.xlsx"):
            with self.subTest(changed_path=changed_path):
                versions = self.tab_apply_versions()

                def change_file():
                    changed_path.write_bytes(changed_path.read_bytes() + b"external change")

                with patch.object(mark_finder.batch, "_restore_tab", wraps=mark_finder.batch._restore_tab) as restore, \
                        patch.object(mark_finder, "_write_format_color", wraps=mark_finder._write_format_color) as write:
                    with self.assertRaisesRegex(ValueError, "发生变化.*重新扫描"):
                        self.tab_apply(SimpleNamespace(ColorIndex=-4142), include_cell=True,
                                       versions=versions, on_open=change_file)
                restore.assert_not_called()
                write.assert_not_called()
                self.last_tab_book.Worksheets.assert_not_called()
                self.last_tab_sheet.Range.assert_not_called()
                self.assertEqual(self.last_tab_sheet.Tab.ColorIndex, 6)
                self.last_tab_book.Save.assert_not_called()
                self.last_tab_book.Close.assert_called_once_with(SaveChanges=False)

    def test_version_changes_after_sheet_writes_abort_before_save(self):
        for changed_path in (self.path, self.root / "old.xlsx"):
            with self.subTest(changed_path=changed_path):
                versions = self.tab_apply_versions()

                def change_file(*_):
                    self.assertEqual(self.last_tab_sheet.Tab.ColorIndex, -4142)
                    self.assertEqual(self.last_tab_sheet.Range.return_value.Interior.Pattern, -4142)
                    changed_path.write_bytes(changed_path.read_bytes() + b"external change")

                with self.assertRaisesRegex(ValueError, "发生变化.*重新扫描"):
                    self.tab_apply(SimpleNamespace(ColorIndex=-4142), include_cell=True,
                                   versions=versions, before_save=change_file)
                self.last_tab_book.Save.assert_not_called()
                self.last_tab_book.Close.assert_called_once_with(SaveChanges=False)

    def test_unchanged_versions_restore_and_save_normally(self):
        results, book, tab, normalize = self.tab_apply(SimpleNamespace(ColorIndex=-4142), include_cell=True,
                                                      versions=self.tab_apply_versions())
        self.assertEqual([item["status"] for item in results], ["ok", "ok"])
        self.assertEqual(tab.ColorIndex, -4142)
        normalize.assert_called_once()
        book.Save.assert_called_once()
        book.Close.assert_called_once_with(SaveChanges=False)

    def test_legacy_apply_without_versions_never_reads_fingerprints(self):
        with patch("core.link_repair.fingerprint", side_effect=AssertionError("legacy API must not read files")) as fingerprint:
            results, book, _, _ = self.tab_apply(SimpleNamespace(ColorIndex=-4142), include_cell=True)
        fingerprint.assert_not_called()
        self.assertEqual([item["status"] for item in results], ["ok", "ok"])
        book.Save.assert_called_once()
        book.Close.assert_called_once_with(SaveChanges=False)

    def test_binary_color_repair_opens_a_copy_and_replaces_only_after_success(self):
        from core.link_repair import fingerprint

        path, backup = self.root / "current.xls", self.root / "backup.xls"
        path.write_bytes(b"current")
        backup.write_bytes(b"backup")
        versions = fingerprint(path), fingerprint(backup)
        working_paths = []

        def restore(excel, working, old_path, cells, check_versions):
            working = Path(working)
            working_paths.append(working)
            self.assertNotEqual(working, path)
            self.assertEqual(working.name, path.name)
            self.assertEqual(working.parent.parent, path.parent)
            working.write_bytes(b"Excel updated metadata and restored colors")
            check_versions()
            return [{"status": "ok"}]

        with patch.object(mark_finder, "_apply_colors", side_effect=restore):
            self.assertEqual(mark_finder.apply_file(Mock(), str(path), str(backup), [], versions=versions), [{"status": "ok"}])
        self.assertEqual(path.read_bytes(), b"Excel updated metadata and restored colors")
        self.assertEqual(fingerprint(backup), versions[1])
        self.assertFalse(working_paths[0].parent.exists())

    def test_binary_color_repair_keeps_external_changes_and_discards_its_copy(self):
        from core.link_repair import fingerprint

        path, backup = self.root / "current.xls", self.root / "backup.xls"
        for changed_path in (path, backup):
            with self.subTest(changed_path=changed_path):
                path.write_bytes(b"current")
                backup.write_bytes(b"backup")
                versions = fingerprint(path), fingerprint(backup)

                def restore(excel, working, old_path, cells, check_versions):
                    Path(working).write_bytes(b"repaired")
                    changed_path.write_bytes(b"external changes")
                    return [{"status": "ok"}]

                with patch.object(mark_finder, "_apply_colors", side_effect=restore), \
                        patch.object(mark_finder.os, "replace") as replace:
                    with self.assertRaisesRegex(ValueError, "发生变化.*重新扫描"):
                        mark_finder.apply_file(Mock(), str(path), str(backup), [], versions=versions)
                replace.assert_not_called()
                self.assertEqual(changed_path.read_bytes(), b"external changes")
                self.assertFalse(list(self.root.glob(".color-repair-*")))

    def test_binary_color_repair_never_replaces_original_on_failure_or_no_changes(self):
        from core.link_repair import fingerprint

        path, backup = self.root / "current.xls", self.root / "backup.xls"
        path.write_bytes(b"current")
        backup.write_bytes(b"backup")
        versions = fingerprint(path), fingerprint(backup)
        for fail in (False, True):
            with self.subTest(fail=fail):
                def restore(excel, working, old_path, cells, check_versions):
                    Path(working).write_bytes(b"Excel updated metadata")
                    if fail:
                        raise RuntimeError("save failed")
                    return [{"status": "failed"}]

                with patch.object(mark_finder, "_apply_colors", side_effect=restore), \
                        patch.object(mark_finder.os, "replace") as replace:
                    if fail:
                        with self.assertRaisesRegex(RuntimeError, "save failed"):
                            mark_finder.apply_file(Mock(), str(path), str(backup), [], versions=versions)
                    else:
                        mark_finder.apply_file(Mock(), str(path), str(backup), [], versions=versions)
                replace.assert_not_called()
                self.assertEqual(fingerprint(path), versions[0])
                self.assertFalse(list(self.root.glob(".color-repair-*")))


def bgr(hex_color):
    return mark_finder._excel_color("#" + hex_color)


def col_name(col):
    return "ABCDEFGHIJ"[col - 1]


class FakeRange:
    """Excel Range 替身：整块颜色一样时 Interior.Color 给颜色，不一样给 None（和真 Excel 一样）"""
    def __init__(self, sheet, row, col, rows=1, cols=1):
        self.sheet, self.row, self.col, self.n_rows, self.n_cols = sheet, row, col, rows, cols
        self.Rows, self.Columns = Mock(Count=rows), Mock(Count=cols)
        sheet.calls += 1

    def _cells(self):
        return [(r, c) for r in range(self.row, self.row + self.n_rows) for c in range(self.col, self.col + self.n_cols)]

    @property
    def Interior(self):
        colors = {self.sheet.fill.get(rc, (bgr("FFFFFF"), -4142))[0] for rc in self._cells()}
        pattern = self.sheet.fill.get((self.row, self.col), (0, -4142))[1]
        return Mock(Color=colors.pop() if len(colors) == 1 else None, Pattern=pattern)

    @property
    def Font(self):
        return Mock(Color=self.sheet.font.get((self.row, self.col), 0))

    def Offset(self, rows, cols):
        return FakeRange(self.sheet, self.row + rows, self.col + cols, self.n_rows, self.n_cols)

    def Resize(self, rows, cols):
        return FakeRange(self.sheet, self.row, self.col, rows, cols)

    @property
    def Cells(self):
        return [FakeRange(self.sheet, r, c) for r, c in self._cells()]

    @property
    def MergeCells(self):
        return any((self.row, self.col) in area for area in self.sheet.merged)

    @property
    def MergeArea(self):
        area = next(a for a in self.sheet.merged if (self.row, self.col) in a)
        first, last = min(area), max(area)
        return Mock(Address=f"${col_name(first[1])}${first[0]}:${col_name(last[1])}${last[0]}")

    @property
    def Address(self):  # pywin32 里 Address 是属性，返回带 $ 的字符串
        return f"${col_name(self.col)}${self.row}"


class FakeSheet:
    def __init__(self, name, rows, cols):
        self.Name, self.rows, self.cols = name, rows, cols
        self.fill, self.font, self.merged, self.calls = {}, {}, [], 0

    def paint(self, row, col, fill, font=None):
        self.fill[(row, col)] = (bgr(fill), 1)
        if font:
            self.font[(row, col)] = bgr(font)

    @property
    def UsedRange(self):
        return FakeRange(self, 1, 1, self.rows, self.cols)

    def Range(self, addr):
        return FakeRange(self, int(addr[1:]), "ABCDEFGHIJ".index(addr[0]) + 1)


class ExcelReadTests(unittest.TestCase):
    """.xls / .xlsb 用后台 Excel 读颜色"""
    def scan(self, sheet, suffix=".xls", saved=None, recorded=(), colors=None):
        book = Mock(Worksheets=[sheet])
        excel = Mock()
        excel.Workbooks.Open.return_value = book
        hits = mark_finder.scan_file_excel(excel, Path("C:/x/book" + suffix), mark_finder._styles(MAGENTA if colors is None else colors),
                                           set(recorded), saved or {})
        self.assertEqual(excel.Workbooks.Open.call_args.kwargs["ReadOnly"], True)
        book.Close.assert_called_once_with(SaveChanges=False)
        return {h["addr"]: h for h in hits}

    def test_marker_cells_are_found_merged_area_once(self):
        sheet = FakeSheet("明細", 40, 8)
        sheet.paint(5, 1, "C0007A", "FFFFFF")      # A5 底色字色都对上
        sheet.paint(7, 2, "C0007A", "123456")      # B7 只有底色
        sheet.paint(9, 3, "C0007B", "FFFFFF")      # 差一点不算
        sheet.paint(3, 4, "FFFF00")                # 别的颜色
        sheet.merged.append({(2, c) for c in (5, 6, 7)})
        for c in (5, 6, 7):
            sheet.paint(2, c, "000000", "FFFF00")  # E2:G2 合并区域算一格
        hits = self.scan(sheet, recorded=[("明細", "A5")])
        self.assertEqual({a: h["match"] for a, h in hits.items()}, {"A5": "both", "B7": "fill", "E2:G2": "both"})
        self.assertEqual(hits["B7"]["font"], "#123456")
        self.assertTrue(hits["A5"]["recorded"])
        self.assertLess(sheet.calls, 40 * 8)  # 对半切开找，不是逐格读

    def test_xls_color_replaced_by_palette_color_still_counts(self):
        sheet = FakeSheet("表", 10, 3)
        sheet.paint(4, 2, "993366", "FFFFFF")      # 存进 .xls 后标记色被换成了 56 色里的这一种
        saved = {("C0007A", "FFFFFF"): ("993366", "FFFFFF")}
        hits = self.scan(sheet, saved=saved)
        self.assertEqual(hits["B4"]["match"], "both")
        self.assertTrue(hits["B4"]["approx"])
        self.assertEqual(self.scan(sheet, ".xlsb", saved=saved), {})  # .xlsb 不会换色

    def test_blank_sheet_is_one_call(self):
        sheet = FakeSheet("空", 500, 10)
        self.assertEqual(self.scan(sheet), {})
        self.assertLessEqual(sheet.calls, 2)

    def test_binary_tab_markers_are_found_without_any_cell_markers(self):
        sheet = FakeSheet("OnlyTab", 1, 1)
        sheet.Tab = SimpleNamespace(ColorIndex=6, Color=mark_finder._excel_color("#FFC000"))
        hits = self.scan(sheet)
        self.assertEqual(hits["A1"]["kind"], "tab")

    def test_history_with_same_fill_and_different_fonts_matches_every_pair(self):
        colors = MAGENTA + [{"fill": "#C0007A", "font": "#123456"}]
        sheet = FakeSheet("表", 10, 3)
        sheet.paint(1, 1, "C0007A", "FFFFFF")
        sheet.paint(2, 1, "C0007A", "123456")
        sheet.paint(3, 1, "C0007A", "000000")
        for suffix in (".xls", ".xlsb"):
            with self.subTest(suffix=suffix):
                hits = self.scan(sheet, suffix, colors=colors)
                self.assertEqual({a: h["match"] for a, h in hits.items()},
                                 {"A1": "both", "A2": "both", "A3": "fill"})

    def test_xls_palette_collisions_preserve_fonts_and_prefer_direct_matches(self):
        colors = MAGENTA + [{"fill": "#993366", "font": "#123456"},
                            {"fill": "#C0007A", "font": "#000000"}]
        saved = {("C0007A", "FFFFFF"): ("993366", "FFFFFF"),
                 ("C0007A", "000000"): ("993366", "000000")}
        sheet = FakeSheet("表", 10, 3)
        for row, font in enumerate(("FFFFFF", "000000", "123456", "654321"), 1):
            sheet.paint(row, 1, "993366", font)
        hits = self.scan(sheet, saved=saved, colors=colors)
        self.assertEqual({a: h["match"] for a, h in hits.items()},
                         {"A1": "both", "A2": "both", "A3": "both", "A4": "fill"})
        self.assertTrue(hits["A1"]["approx"])
        self.assertTrue(hits["A2"]["approx"])
        self.assertNotIn("approx", hits["A3"])
        self.assertNotIn("approx", hits["A4"])


if __name__ == "__main__":
    unittest.main()
