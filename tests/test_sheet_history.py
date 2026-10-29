"""删除同名旧表后重命名新版、连续改名、备份阶段和严格内容验证。"""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.styles import PatternFill

from core import color_compare, hidden_repair, link_repair, mark_finder, sheet_history as history, workflow
from core.backup_compare import pair_sheets


def event(before, after, deleted=(), renamed=None, frozen=(), before_hash="v0", after_hash="v1"):
    return {"at": "2026-10-28T12:00:00+09:00", "before": before, "after": after,
            "deleted": list(deleted), "renamed": renamed or {}, "preserved_values": list(frozen),
            "before_hash": before_hash, "after_hash": after_hash}


class SheetHistoryTests(unittest.TestCase):
    def setUp(self):
        self.events = [event(["Layout", "新Layout", "Keep"], ["Layout", "Keep"],
                             ["Layout"], {"新Layout": "Layout"})]

    def match(self, projection, current, old):
        return pair_sheets(current, old, title=lambda n: n, code=lambda n: "", words=lambda n: set(), history=projection)

    def test_history_precedes_same_name_and_deleted_sheet_is_never_selected(self):
        projected = history.project(self.events, ["Layout", "Keep"], self.events[0]["before"])
        self.assertEqual(self.match(projected, ["Layout", "Keep"], self.events[0]["before"]),
                         {"Layout": ("新Layout", "history"), "Keep": ("Keep", "history")})

    def test_chart_sheet_in_workbook_history_does_not_prevent_worksheet_pairing(self):
        self.events[0]["before"].append("ChartPage")
        self.events[0]["after"].append("ChartPage")
        projected = history.project(self.events, ["Layout", "Keep", "ChartPage"], self.events[0]["before"])
        matched = self.match(projected, ["Layout", "Keep"], ["Layout", "新Layout", "Keep"])
        self.assertEqual(matched["Layout"], ("新Layout", "history"))

    def test_multiple_renames_trace_back_to_the_selected_backup_stage(self):
        self.events.append(event(["Layout", "Keep"], ["Final", "Keep"], renamed={"Layout": "Final"},
                                 before_hash="v1", after_hash="v2"))
        original = history.project(self.events, ["Final", "Keep"], self.events[0]["before"])
        middle = history.project(self.events, ["Final", "Keep"], ["Layout", "Keep"])
        self.assertEqual(original.origins["Final"], "新Layout")
        self.assertEqual(middle.origins["Final"], "Layout")
        self.assertEqual(len(middle.operations), 1)

    def test_reused_names_require_backup_hash_when_history_stages_disagree(self):
        events = [event(["A", "B"], ["X", "B"], renamed={"A": "X"}),
                  event(["X", "B"], ["X", "A"], renamed={"B": "A"}, before_hash="v1", after_hash="v2"),
                  event(["X", "A"], ["B", "A"], renamed={"X": "B"}, before_hash="v2", after_hash="v3")]
        uncertain = history.project(events, ["A", "B"], ["A", "B"])
        self.assertEqual(uncertain.origins, {"A": None, "B": None})
        self.assertEqual(self.match(uncertain, ["A", "B"], ["A", "B"]), {})
        exact = history.project(events, ["A", "B"], ["A", "B"], "v0")
        self.assertEqual(exact.origins, {"A": "B", "B": "A"})

    def test_history_is_scoped_to_current_structure_and_validated_before_import(self):
        self.assertFalse(history.project(self.events, ["Other"], self.events[0]["before"]).origins)
        records = {"path": self.events}
        self.assertTrue(history.validate(records))
        invalid = copy.deepcopy(records)
        invalid["path"][0]["renamed"]["Missing"] = "Fake"
        self.assertFalse(history.validate(invalid))

    def test_reorder_history_validates_and_preserves_identity_across_rename(self):
        events = [event(["A", "B", "C"], ["C", "A", "B"], before_hash="v0", after_hash="v1"),
                  event(["C", "A", "B"], ["B", "C", "Final"], renamed={"A": "Final"},
                        before_hash="v1", after_hash="v2")]
        self.assertTrue(history.validate({"path": events}))
        projected = history.project(events, ["B", "C", "Final"], ["A", "B", "C"], "v0")
        self.assertEqual(projected.origins, {"B": "B", "C": "C", "Final": "A"})
        invalid = copy.deepcopy(events)
        invalid[0]["after"] = ["C", "A", "Other"]
        self.assertFalse(history.validate({"path": invalid}))

    def test_only_recorded_formula_effects_are_projected_and_strings_and_external_refs_stay(self):
        frozen = {"sheet": "新Layout", "addr": "A1", "formula": "=Layout!B1", "value": 42, "kind": "cell"}
        self.events[0]["preserved_values"] = [frozen]
        p = history.project(self.events, ["Layout", "Keep"], self.events[0]["before"])
        self.assertEqual(p.value("新Layout", "A1", "=Layout!B1"), 42)
        self.assertNotEqual(p.value("新Layout", "A2", "=Layout!B1"), 42)
        self.assertNotEqual(p.value("新Layout", "A1", "=Layout!C1"), 42)
        self.assertEqual(p.value("Keep", "A1", "='新Layout'!A1"), history.formula_value("=Layout!A1"))
        expression = '=IF(A1="新Layout!A1",\'[other.xlsx]新Layout\'!A1,\'新Layout\'!A1)'
        projected = p.value("Keep", "A1", expression)
        self.assertIn('"新Layout!A1"', projected)
        self.assertIn("'[other.xlsx]新Layout'!A1", projected)
        self.assertIn("'layout'!A1", projected)

    def test_hidden_repair_uses_history_even_for_blank_sheet_and_still_rejects_content_changes(self):
        def sheet(name, hidden=False, signature="same"):
            return {"name": name, "code": name, "words": [], "signature": signature, "populated": 0,
                    "visible": -1, "content_error": "", "columns_error": "", "protected": False,
                    "columns": [{"first": 1, "last": 10, "hidden": hidden, "width": None if hidden else 10}]}

        current = {"sheets": [sheet("Layout"), sheet("Keep")], "visible_count": 2, "structure_protected": False}
        old = {"sheets": [sheet("Layout", signature="wrong"), sheet("新Layout", hidden=True), sheet("Keep")]}
        projected = history.project(self.events, ["Layout", "Keep"], self.events[0]["before"])
        planned = hidden_repair.plan_changes(current, old, hidden_repair.DEFAULT_OPTIONS, projected)
        self.assertEqual([(i["old_sheet"], i["action"]) for i in planned], [("新Layout", "restore")])
        current["sheets"][0]["signature"] = "moved cells"
        self.assertEqual(hidden_repair.plan_changes(current, old, hidden_repair.DEFAULT_OPTIONS, projected)[0]["action"], "skip")

    def test_link_pairing_uses_the_same_history(self):
        info = lambda names: [{"name": n, "code": n, "words": []} for n in names]
        projected = history.project(self.events, ["Layout", "Keep"], self.events[0]["before"])
        pairs = link_repair._pair_sheets(info(["Layout", "Keep"]), info(self.events[0]["before"]), projected)
        self.assertEqual(pairs["Layout"][0]["name"], "新Layout")

    def test_content_signature_accepts_only_known_formula_changes_and_keeps_coordinates(self):
        from types import SimpleNamespace

        def sample(name, values, row=1):
            return SimpleNamespace(Name=name, UsedRange=SimpleNamespace(Row=row, Column=1,
                                   Rows=SimpleNamespace(Count=1), Columns=SimpleNamespace(Count=len(values))),
                                   Cells=lambda r, c: (r, c), Range=lambda *args: SimpleNamespace(Formula=(tuple(values),)))

        p = history.project(self.events, ["Layout", "Keep"], self.events[0]["before"])
        old = hidden_repair.content_info(sample("Keep", ["=新Layout!A1", "stable"]), p)["signature"]
        self.assertEqual(old, hidden_repair.content_info(sample("Keep", ["=Layout!A1", "stable"]))["signature"])
        self.assertNotEqual(old, hidden_repair.content_info(sample("Keep", ["=Layout!B1", "stable"]))["signature"])
        self.assertNotEqual(old, hidden_repair.content_info(sample("Keep", ["=Layout!A1", "stable"], row=2))["signature"])

    def test_content_signature_uses_typed_constant_after_formula_is_frozen_by_excel(self):
        from types import SimpleNamespace

        def sample(formula, value):
            return SimpleNamespace(Name="Keep", UsedRange=SimpleNamespace(Row=1, Column=1,
                                   Rows=SimpleNamespace(Count=1), Columns=SimpleNamespace(Count=1)),
                                   Cells=lambda r, c: (r, c), Range=lambda *args: SimpleNamespace(Formula=formula, Value2=value))

        operation = event(["Drop", "Keep"], ["Keep"], ["Drop"], frozen=[
            {"sheet": "Keep", "addr": "A1", "formula": "=Drop!A1", "value": 42.0, "kind": "cell"}])
        projection = history.project([operation], ["Keep"], ["Drop", "Keep"])
        original = hidden_repair.content_info(sample("=Drop!A1", 42.0), projection)["signature"]
        self.assertEqual(original, hidden_repair.content_info(sample("42", 42.0))["signature"])
        self.assertNotEqual(original, hidden_repair.content_info(sample("42", "42"))["signature"])


class SheetHistoryFileTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.path, self.old = self.root / "current.xlsx", self.root / "old.xlsx"
        patcher = patch.object(workflow, "STATE", {"sheet_history": {}})
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_color_plan_uses_original_sheet_colors_and_accounts_for_formula_renames_and_frozen_values(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        old.active.title = "Layout"
        old.create_sheet("新Layout")
        old.create_sheet("Keep")
        old["Layout"]["A1"] = "wrong old version"
        old["新Layout"]["A1"] = "=Layout!A2"
        old["新Layout"]["A2"] = "=Keep!A1"
        old["新Layout"]["A1"].fill = PatternFill("solid", fgColor="112233")
        old["新Layout"]["A2"].fill = PatternFill("solid", fgColor="223344")
        old["Keep"]["A1"] = "=新Layout!B1"
        old.save(self.old)
        current.active.title = "Layout"
        current.create_sheet("Keep")
        current["Layout"]["A1"] = 42
        current["Layout"]["A2"] = "='Keep'!A1"
        current["Keep"]["A1"] = "=Layout!B1"
        for addr in ("A1", "A2"):
            current["Layout"][addr].fill = PatternFill("solid", fgColor="FF0000")
        current.save(self.path)
        history.record(self.path, old.sheetnames, current.sheetnames, ["Layout"], {"新Layout": "Layout"},
                       [{"sheet": "新Layout", "addr": "A1", "formula": "=Layout!A2", "value": 42, "kind": "cell"}],
                       history.fingerprint(self.old))
        planned = color_compare.plan_file(None, self.path, self.old)["cells"]
        self.assertEqual([(i["old_sheet"], i["action"]) for i in planned], [("新Layout", "restore"), ("新Layout", "restore")])
        hits = [{"sheet": "Layout", "addr": "A1", "fill": "#FF0000", "font": ""}]
        self.assertEqual(mark_finder.plan_file(self.path, self.old, hits)[0]["action"], "restore")
        current["Layout"]["A1"] = 99
        current.save(self.path)
        self.assertEqual(color_compare.plan_file(None, self.path, self.old)["cells"][0]["action"], "skip")
        old.close()
        current.close()

    def test_file_rename_moves_history_without_mixing_same_basename_in_other_folder(self):
        self.path.write_bytes(b"saved")
        history.record(self.path, ["Old"], ["New"], [], {"Old": "New"}, [], "before")
        other = self.root / "other" / self.path.name
        self.assertEqual(history.get(other), [])
        renamed = self.root / "★current.xlsx"
        history.move(self.path, renamed)
        self.assertEqual(history.get(self.path), [])
        self.assertEqual(history.get(renamed)[0]["renamed"], {"Old": "New"})


if __name__ == "__main__":
    unittest.main()
