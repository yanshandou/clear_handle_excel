"""没有清理历史时的 Sheet 推断与候选文件自动写入保护。"""
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import openpyxl
from openpyxl.styles import PatternFill

from core import color_compare, hidden_repair, link_repair, mark_finder, sheet_history, workflow
from core.backup_compare import inferred_renames, pair_sheets


def s(name, code="", words=()):
    return {"name": name, "code": code, "words": set(words)}


def matches(current, old, history=None):
    return pair_sheets(current, old, title=lambda x: x["name"], code=lambda x: x["code"],
                       words=lambda x: x["words"], history=history)


class SheetGuessTests(unittest.TestCase):
    def test_new_version_can_win_against_same_name_by_unique_code_or_content(self):
        old = [s("Layout", "Old", ["old"]), s("新Layout", "New", ["new", "title"])]
        for current in (s("Layout", "New"), s("Layout", "", ["new", "title"])):
            with self.subTest(current=current):
                result = matches([current], old)
                self.assertEqual(result["Layout"][0]["name"], "新Layout")
                self.assertEqual(result["Layout"][1], "rule")
                self.assertEqual(inferred_renames(result, lambda x: x["name"]), {"新Layout": "Layout"})

    def test_same_name_can_win_when_evidence_points_to_it(self):
        old = [s("Layout", "Original", ["title"]), s("新Layout", "New", ["new"])]
        self.assertEqual(matches([s("Layout", "Original", ["title"])], old)["Layout"][0]["name"], "Layout")

    def test_no_evidence_or_contradictory_evidence_never_falls_back_to_same_name(self):
        old = [s("Layout", "Old", ["original", "title"]), s("新Layout", "New", ["new", "title"])]
        self.assertEqual(matches([s("Layout")], old), {})
        self.assertEqual(matches([s("Layout", "Old", ["new", "title"])], old), {})
        self.assertEqual(matches([s("Layout", words=["new"])], old), {})
        self.assertEqual(matches([s("Layout", "", ["shared"])], [s("Layout", words=["shared"]), s("新Layout", words=["shared"])]), {})

    def test_unchanged_new_sheet_is_not_stolen_and_reserved_words_are_not_prefixes(self):
        old = [s("Layout", "A"), s("新Layout", "B")]
        result = matches([s("Layout", "A"), s("新Layout", "B")], old)
        self.assertEqual(result["Layout"][0]["name"], "Layout")
        self.assertEqual(result["新Layout"][0]["name"], "新Layout")
        result = matches([s("規約")], [s("規約"), s("新規約")])
        self.assertEqual(result["規約"][0]["name"], "規約")

    def test_historical_ambiguity_cannot_be_overridden_by_guessing(self):
        old = [s("Layout", "Old"), s("新Layout", "New")]
        self.assertEqual(matches([s("Layout", "New")], old, sheet_history.Projection({"Layout": None})), {})

    def test_formula_inference_does_not_treat_reused_old_sheet_references_as_safe(self):
        projection = sheet_history.with_inferred_renames(sheet_history.Projection(), {"新Layout": "Layout"})
        self.assertEqual(projection.value("新Layout", "A1", "=新Layout!B1"), sheet_history.formula_value("=Layout!B1"))
        self.assertNotEqual(projection.value("新Layout", "A1", "=Layout!B1"), sheet_history.formula_value("=Layout!B1"))
        self.assertNotEqual(projection.value("新Layout", "A1", '=INDIRECT("Layout!B1")'), '=INDIRECT("Layout!B1")')


class ColorGuessTests(unittest.TestCase):
    def test_color_comparison_uses_inferred_new_sheet_and_checks_formula_identity(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(workflow, "STATE", {"sheet_history": {}}):
            root = Path(directory)
            current_path, old_path = root / "current.xlsx", root / "old.xlsx"
            old, current = openpyxl.Workbook(), openpyxl.Workbook()
            old.active.title = "Layout"
            old.active["A1"] = "unrelated"
            new = old.create_sheet("新Layout")
            new["A1"] = "title"
            new["A2"] = "=新Layout!A1"
            new["A3"] = "=Layout!A1"
            for row in range(1, 4):
                new.cell(row, 1).fill = PatternFill("solid", fgColor="112233")
            current.active.title = "Layout"
            current.active["A1"] = "title"
            current.active["A2"] = "=Layout!A1"
            current.active["A3"] = "=Layout!A1"
            for row in range(1, 4):
                current.active.cell(row, 1).fill = PatternFill("solid", fgColor="FF0000")
            old.save(old_path)
            current.save(current_path)
            plan = color_compare.plan_file(None, current_path, old_path)["cells"]
            by_addr = {item["addr"]: item for item in plan}
            self.assertEqual([by_addr[a]["old_sheet"] for a in ("A1", "A2", "A3")], ["新Layout"] * 3)
            self.assertEqual([by_addr[a]["action"] for a in ("A1", "A2", "A3")], ["restore", "restore", "skip"])
            old.close()
            current.close()


class FileGuessSafetyTests(unittest.TestCase):
    def test_inferred_backup_is_shown_but_never_automatically_applied(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            current, old = root / "current", root / "old"
            current.mkdir()
            old.mkdir()
            path, backup = current / "book.xlsx", old / "book_backup.xlsx"
            path.write_bytes(b"current")
            backup.write_bytes(b"old")
            plan = {"cells": [{"sheet": "Main", "addr": "A1", "kind": "sheet", "action": "restore", "reason": "ready"}]}
            state = {"busy": True, "maintenance": "repair"}
            com = SimpleNamespace(CoInitialize=Mock(), CoUninitialize=Mock())
            with patch.object(workflow, "STATE", state), patch.object(workflow, "pythoncom", com), \
                 patch.object(workflow, "log_msg"), patch.object(hidden_repair, "start_background_excel", return_value=Mock()), \
                 patch.object(hidden_repair, "quit_background_excel"), patch.object(hidden_repair, "plan_file", return_value=plan), \
                 patch.object(hidden_repair, "apply_file") as apply, patch.object(hidden_repair, "_STATUS", {}):
                hidden_repair._run(current, [path], old, [backup], hidden_repair.DEFAULT_OPTIONS)
                entry = hidden_repair.status()["applied"]["files"][0]
                self.assertEqual(entry["old"], str(backup))
                self.assertIn("人工确认", entry["error"])
                self.assertEqual(entry["cells"][0]["action"], "skip")
                apply.assert_not_called()

            state.update(busy=True, maintenance="repair")
            hit = {"sheet": "Main", "addr": "A1", "formula": "=#REF!"}
            with patch.object(workflow, "STATE", state), patch.object(workflow, "pythoncom", com), \
                 patch.object(workflow, "log_msg"), patch.object(link_repair, "start_background_excel", return_value=Mock()), \
                 patch.object(link_repair, "quit_background_excel"), patch.object(link_repair, "scan_file", return_value={"hits": [hit]}), \
                 patch.object(link_repair, "preview_file", return_value={"cells": [{**hit, "action": "restore"}]}), \
                 patch.object(link_repair, "apply_file") as apply, patch.object(link_repair, "_STATUS", {}):
                link_repair._run("repair", current, [path], old, [backup])
                entry = link_repair.status()["applied"]["files"][0]
                self.assertEqual(entry["old"], str(backup))
                self.assertEqual(entry["cells"][0]["action"], "skip")
                apply.assert_not_called()

            state.update(busy=True, maintenance="repair")
            with patch.object(workflow, "STATE", state), patch.object(workflow, "pythoncom", com), \
                 patch.object(workflow, "log_msg"), patch("core.excel_ops.quit_background_excel"), \
                 patch.object(mark_finder, "_clear_mark_files", return_value=[]), \
                 patch.object(mark_finder, "compare_file_colors", return_value={"cells": [{**hit, "action": "restore"}]}), \
                 patch.object(mark_finder, "apply_file") as apply, patch.object(mark_finder, "_STATUS", {}):
                mark_finder._run_repair(current, [path], old, [backup], [])
                entry = mark_finder.status()["applied"]["files"][0]
                self.assertEqual(entry["old"], str(backup))
                self.assertEqual(entry["cells"][0]["action"], "skip")
                apply.assert_not_called()


if __name__ == "__main__":
    unittest.main()
