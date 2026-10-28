import tempfile
import unittest
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Color, Font, GradientFill, PatternFill

from core import color_compare


class AllColorComparisonTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.path, self.old = self.root / "current.xlsx", self.root / "backup.xlsx"

    def plan(self, current, old):
        old.save(self.old)
        current.save(self.path)
        old.close()
        current.close()
        return color_compare.plan_file(None, self.path, self.old)["cells"]

    def test_arbitrary_fill_font_only_and_missing_original_colors_are_compared(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        for book in (old, current):
            for addr in ("A1", "A2", "A3", "A4"):
                book.active[addr] = addr
        source, sheet = old.active, current.active
        source["A1"].fill = PatternFill("solid", fgColor="123456")
        sheet["A1"].fill = PatternFill("solid", fgColor="8ACDD1")
        source["A2"].font = Font(color="0055AA", bold=True, size=15)
        sheet["A2"].font = Font(color="FE019A", italic=True, size=11)
        source["A3"].fill = copy(sheet["A3"].fill)
        source["A3"].font = Font(color="112233", bold=True)
        sheet["A3"].font = Font(color="112233", bold=False)
        source["A4"].fill = PatternFill("solid", fgColor="AABBCC")
        rows = {item["addr"]:item for item in self.plan(current, old)}
        self.assertEqual(set(rows), {"A1", "A2", "A4"})
        self.assertTrue(all(item["mode"] == "all" and item["action"] == "restore" for item in rows.values()))
        self.assertEqual((rows["A1"]["fill"], rows["A1"]["old_fill"]), ("#8ACDD1", "#123456"))
        self.assertEqual((rows["A2"]["font"], rows["A2"]["old_font"]), ("#FE019A", "#0055AA"))
        self.assertEqual(rows["A4"]["fill"], "none")

    def test_all_tab_colors_and_removed_tab_colors_compare_and_unchanged_colors_stay(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        for book in (old, current):
            book.active.title = "Changed"
            book.create_sheet("Removed")
            book.create_sheet("Same")
        old["Changed"].sheet_properties.tabColor = "112233"
        current["Changed"].sheet_properties.tabColor = "FF0000"
        old["Removed"].sheet_properties.tabColor = Color(theme=4, tint=0.4)
        old["Same"].sheet_properties.tabColor = current["Same"].sheet_properties.tabColor = "C0007A"
        rows = {item["sheet"]:item for item in self.plan(current, old)}
        self.assertEqual(set(rows), {"Changed", "Removed"})
        self.assertEqual(rows["Changed"]["fill"], "#FF0000")
        self.assertEqual(rows["Removed"]["fill"], "none")

    def test_equivalent_indexed_rgb_and_automatic_font_colors_are_not_differences(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        old.active["A1"] = current.active["A1"] = "same"
        old.active["A1"].fill = PatternFill("solid", fgColor=Color(indexed=10))
        current.active["A1"].fill = PatternFill("solid", fgColor="FF0000")
        old.active["A1"].font = Font()
        current.active["A1"].font = Font(color=Color(auto=True))
        self.assertEqual(self.plan(current, old), [])

    def test_theme_slots_match_equivalent_rgb_for_fill_font_and_tabs(self):
        for index, rgb in enumerate(("FFFFFF", "000000", "EEECE1", "1F497D")):
            with self.subTest(theme=index):
                old, current = openpyxl.Workbook(), openpyxl.Workbook()
                old.active["A1"] = current.active["A1"] = "same"
                old.active["A1"].fill = PatternFill("solid", fgColor=Color(theme=index))
                old.active["A1"].font = Font(color=Color(theme=index))
                old.active.sheet_properties.tabColor = Color(theme=index)
                current.active["A1"].fill = PatternFill("solid", fgColor=rgb)
                current.active["A1"].font = Font(color=rgb)
                current.active.sheet_properties.tabColor = rgb
                self.assertEqual(self.plan(current, old), [])

    def test_theme_slot_lookup_is_independent_of_xml_order(self):
        from xml.etree import ElementTree as ET
        from openpyxl.writer.theme import theme_xml

        book = openpyxl.Workbook()
        root = ET.fromstring(theme_xml)
        scheme = root.find(".//{*}clrScheme")
        scheme[:] = list(reversed(scheme))
        book.loaded_theme = ET.tostring(root)
        context = color_compare._context(book)
        self.assertEqual(context["theme"][:4], ["FFFFFF", "000000", "EEECE1", "1F497D"])
        self.assertEqual(color_compare._color(Color(theme=4), context), "4F81BD")
        book.close()

    def test_theme_black_against_rgb_white_is_a_color_difference(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        old.active["A1"] = current.active["A1"] = "same"
        old.active["A1"].fill = PatternFill("solid", fgColor=Color(theme=1))
        old.active["A1"].font = Font(color=Color(theme=1))
        old.active.sheet_properties.tabColor = Color(theme=1)
        current.active["A1"].fill = PatternFill("solid", fgColor="FFFFFF")
        current.active["A1"].font = Font(color="FFFFFF")
        current.active.sheet_properties.tabColor = "FFFFFF"
        rows = self.plan(current, old)
        self.assertEqual(len(rows), 2)
        cell = next(item for item in rows if item.get("kind") != "tab")
        self.assertEqual((cell["old_fill"], cell["old_font"]), ("#000000", "#000000"))
        self.assertTrue(all(item["action"] == "restore" for item in rows))

    def test_tinted_colors_match_excel_rendered_rgb(self):
        samples = (
            ("123456", 0.4, "3585D5"), ("123456", -0.4, "0B1F33"),
            ("123456", 0.8, "BAD6F1"), ("123456", -0.15, "0F2C48"),
            ("123456", 0.5, "5699DC"), ("123456", 0.5001, "589ADC"),
            ("123456", 0.01, "133659"), ("4F81BD", 0.4, "95B3D7"),
            ("4F81BD", -0.4, "2C4D74"), ("4F81BD", 0.8, "DCE6F1"),
            ("4F81BD", -0.15, "3E6CA6"), ("4F81BD", 0.33333, "89ABD3"),
            ("AA2288", 0.4, "E065C1"), ("AA2288", 0.5, "E680CC"),
            ("FF0000", 0.499984740745262, "FF8080"),
            ("FF0000", 0.7999816888943144, "FFCCCC"),
            ("000000", 0.5, "808080"), ("000000", 0.5001, "818181"),
            ("FFFFFF", -0.1499984740745262, "D9D9D9"),
            ("123456", 0.0, "123456"), ("123456", 0.00001, "123456"),
            ("123456", -1.0, "000000"), ("123456", 1.0, "FFFFFF"),
        )
        context = color_compare._context(openpyxl.Workbook())
        for rgb, tint, expected in samples:
            with self.subTest(rgb=rgb, tint=tint):
                self.assertEqual(color_compare._color(Color(rgb=rgb, tint=tint), context), expected)

    def test_tinted_rgb_and_theme_colors_equal_their_saved_excel_rgb(self):
        for source_color, rgb in ((Color(rgb="123456", tint=0.4), "3585D5"),
                                  (Color(theme=4, tint=0.4), "95B3D7"),
                                  (Color(theme=1, tint=0.499984740745262), "808080")):
            with self.subTest(color=source_color):
                old, current = openpyxl.Workbook(), openpyxl.Workbook()
                old.active["A1"] = current.active["A1"] = "same"
                old.active["A1"].fill = PatternFill("solid", fgColor=source_color)
                old.active["A1"].font = Font(color=source_color)
                old.active.sheet_properties.tabColor = source_color
                current.active["A1"].fill = PatternFill("solid", fgColor=rgb)
                current.active["A1"].font = Font(color=rgb)
                current.active.sheet_properties.tabColor = rgb
                self.assertEqual(self.plan(current, old), [])

    def test_changed_content_formula_merged_area_and_empty_row_are_manual(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        source, sheet = old.active, current.active
        source["A1"], sheet["A1"] = "old", "changed"
        source["A2"], sheet["A2"] = "=1+2", "=2+3"
        source["A3"], sheet["A3"] = "old row", "changed row"
        source.merge_cells("A4:B4")
        source["A4"] = sheet["A4"] = "merged"
        for addr in ("A1", "A2", "B3", "A4"):
            sheet[addr].fill = PatternFill("solid", fgColor="AA2288")
        rows = {item["anchor"]:item for item in self.plan(current, old)}
        for addr in ("A1", "A2", "B3", "A4"):
            self.assertEqual(rows[addr]["action"], "skip")
            self.assertTrue(rows[addr]["reason"])

    def test_new_colored_sheet_and_gradients_are_reported_without_overwriting(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        source, sheet = old.active, current.active
        source["A1"] = sheet["A1"] = "gradient"
        source["A1"].fill = GradientFill(stop=("FF0000", "0000FF"))
        sheet["A1"].fill = GradientFill(stop=("FFFF00", "00FF00"))
        extra = current.create_sheet("Unknown")
        extra.sheet_properties.tabColor = "445566"
        extra["A1"].fill = PatternFill("solid", fgColor="778899")
        rows = self.plan(current, old)
        self.assertEqual(len(rows), 3)
        self.assertTrue(all(item["action"] == "skip" for item in rows))
        self.assertTrue(any("渐变" in item["reason"] for item in rows))

    def test_all_color_comparison_does_not_stop_at_the_old_500_marker_limit(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        for row in range(1, 602):
            old.active.cell(row, 1, "same").fill = PatternFill("solid", fgColor="112233")
            current.active.cell(row, 1, "same").fill = PatternFill("solid", fgColor="446688")
        self.assertEqual(len(self.plan(current, old)), 601)

    def test_row_and_column_default_colors_are_reported_for_manual_comparison(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        old.active.column_dimensions["B"].fill = PatternFill("solid", fgColor="112233")
        current.active.column_dimensions["B"].fill = PatternFill("solid", fgColor="ABCDEF")
        old.active.row_dimensions[3].font = Font(color="123456")
        current.active.row_dimensions[3].font = Font(color="987654")
        rows = self.plan(current, old)
        self.assertEqual({item["kind"] for item in rows}, {"row_color", "column_color"})
        self.assertTrue(all(item["action"] == "skip" for item in rows))

    def test_new_merged_area_does_not_hide_an_internal_original_color_difference(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        old.active["A1"] = current.active["A1"] = "title"
        old.active["A1"].fill = PatternFill("solid", fgColor="112233")
        old.active["B1"].fill = PatternFill("solid", fgColor="AABBCC")
        current.active.merge_cells("A1:B1")
        current.active["A1"].fill = PatternFill("solid", fgColor="112233")
        rows = self.plan(current, old)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["addr"], "B1")
        self.assertEqual(rows[0]["action"], "skip")
        self.assertIn("合并", rows[0]["reason"])


    def test_plan_items_carry_stable_category_codes(self):
        old, current = openpyxl.Workbook(), openpyxl.Workbook()
        source, sheet = old.active, current.active
        source["A1"], sheet["A1"] = "old", "changed"
        source["A1"].fill = PatternFill("solid", fgColor="112233")
        sheet["A1"].fill = PatternFill("solid", fgColor="445566")
        source["A2"] = sheet["A2"] = "gradient"
        source["A2"].fill = GradientFill(stop=("FF0000", "0000FF"))
        sheet["A2"].fill = GradientFill(stop=("FFFF00", "00FF00"))
        source.sheet_properties.tabColor = "112233"
        sheet.sheet_properties.tabColor = "445566"
        source.column_dimensions["B"].fill = PatternFill("solid", fgColor="112233")
        sheet.column_dimensions["B"].fill = PatternFill("solid", fgColor="ABCDEF")
        codes = {item["code"] for item in self.plan(current, old)}
        self.assertTrue({"content_mismatch", "gradient", "pending", "dimension_default"} <= codes)


if __name__ == "__main__":
    unittest.main()
