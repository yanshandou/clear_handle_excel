"""真实 Excel 一键任务验证，仅修改临时样本：python tests/verify_repair_tasks_excel.py。"""
import json
import shutil
import sys
import tempfile
import time
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import link_repair, mark_finder, workflow as wf
from core.excel_ops import pythoncom, quit_background_excel, start_background_excel


def wait_for(module):
    deadline = time.monotonic() + 90
    while module.status().get("running"):
        if time.monotonic() > deadline:
            raise TimeoutError("一键任务超过 90 秒")
        time.sleep(0.2)
    data = module.status()["applied"]
    assert not any(file["error"] for file in data["files"]), data
    assert not wf.STATE["busy"] and "maintenance" not in wf.STATE
    return data


def main():
    with tempfile.TemporaryDirectory(prefix="one_click_repair_") as directory:
        root = Path(directory)
        current, old = root / "current" / "001", root / "backup" / "001"
        current.mkdir(parents=True)
        old.mkdir(parents=True)
        pythoncom.CoInitialize()
        excel = book = None
        try:
            excel = start_background_excel()
            book = excel.Workbooks.Add()
            main_sheet = book.Worksheets(1)
            main_sheet.Name = "Main"
            while book.Worksheets.Count > 1:
                book.Worksheets(book.Worksheets.Count).Delete()
            source = book.Worksheets.Add(After=main_sheet)
            source.Name = "Source"
            source.Range("A1").Value2 = 10
            main_sheet.Range("A1").Formula = "=Source!A1*2"
            main_sheet.Range("B1").Value2 = "original content"
            main_sheet.Range("A2").Formula = "=A1+1"
            main_sheet.Range("A3").Formula = "=#REF!"  # 备份本身也损坏的格子必须保留。
            main_sheet.Range("A1").Interior.Color = 123456
            main_sheet.Range("A1").NumberFormat = "0.00"
            excel.CalculateFull()
            book.SaveAs(str(old / "links.xlsx"), FileFormat=51)
            book.SaveCopyAs(str(current / "links.xlsx"))
            book.Close(SaveChanges=False)
            book = excel.Workbooks.Open(str(current / "links.xlsx"), UpdateLinks=0)
            book.Worksheets("Source").Delete()
            book.Save()
            book.Close(SaveChanges=False)
            book = None
        finally:
            if book is not None:
                book.Close(SaveChanges=False)
            quit_background_excel(excel)
            pythoncom.CoUninitialize()

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Main"
        for addr, text in (("A1", "restore"), ("A2", "original marker"), ("A3", "old content"), ("B4", "merged")):
            sheet[addr] = text
            sheet[addr].font = Font(color="112233")
            sheet[addr].fill = PatternFill("solid", fgColor="DDEEFF")
        sheet.merge_cells("B4:C4")
        sheet["A2"].fill = PatternFill("solid", fgColor="C0007A")
        book.save(old / "colors.xlsx")
        book.close()
        shutil.copy2(old / "colors.xlsx", current / "colors.xlsx")
        book = openpyxl.load_workbook(current / "colors.xlsx")
        sheet = book["Main"]
        for addr in ("A1", "A2", "A3", "B4"):
            sheet[addr].fill = PatternFill("solid", fgColor="C0007A")
            sheet[addr].font = Font(color="FFFFFF")
        sheet["A3"] = "changed content"
        book.save(current / "colors.xlsx")
        book.close()

        stamps = {path: link_repair.fingerprint(path) for path in old.iterdir()}
        with patch.object(wf, "STATE", {"busy": False, "job": None}), patch.object(wf, "log_msg"):
            link_repair.repair(str(current), str(old.parent))
            links = wait_for(link_repair)
            mark_finder.repair(str(current), str(old.parent), [{"fill": "#C0007A", "font": "#FFFFFF"}])
            colors = wait_for(mark_finder)

        book = openpyxl.load_workbook(current / "links.xlsx", data_only=False)
        sheet = book["Main"]
        assert sheet["A1"].value == 20 and sheet["A1"].number_format == "0.00"
        assert sheet["A2"].value == "=A1+1" and sheet["A3"].value == "=#REF!"
        assert sheet["A1"].fill.fgColor.rgb[-6:] == "40E201"
        book.close()
        book = openpyxl.load_workbook(current / "colors.xlsx")
        sheet = book["Main"]
        assert sheet["A1"].fill.fgColor.rgb[-6:] == "DDEEFF" and sheet["A1"].font.color.rgb[-6:] == "112233"
        assert sheet["B4"].fill.fgColor.rgb[-6:] == "DDEEFF" and "B4:C4" in sheet.merged_cells
        assert sheet["A2"].fill.fgColor.rgb[-6:] == "C0007A" and sheet["A3"].value == "changed content"
        book.close()
        assert all(link_repair.fingerprint(path) == stamp for path, stamp in stamps.items())
        link_cells = next(file["cells"] for file in links["files"] if file["name"] == "links.xlsx")
        color_cells = next(file["cells"] for file in colors["files"] if file["name"] == "colors.xlsx")
        assert any(cell.get("action") == "skip" for cell in link_cells)
        assert sum(cell.get("status") == "ok" for cell in color_cells) == 2
        assert sum(cell.get("action") == "skip" for cell in color_cells) == 2
        print(json.dumps({"one_click_links": True, "one_click_colors": True, "skipped_details_preserved": True,
                          "merged_cells_preserved": True, "backup_unchanged": True, "background_lock_released": True}))


if __name__ == "__main__":
    main()
