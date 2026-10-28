"""真实 Excel 标签色恢复验证，仅创建临时样本：python tests/verify_tab_color_repair_excel.py。"""
import json
import sys
import tempfile
import time
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.styles import Color, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import batch, color_compare, mark_finder as repair, workflow as wf
from core.excel_ops import pythoncom, quit_background_excel, start_background_excel
from core.link_repair import fingerprint


def create_sample(excel, current, old, name, file_format):
    book = excel.Workbooks.Add()
    try:
        book.Worksheets(1).Name = "NoColor"
        while book.Worksheets.Count > 1:
            book.Worksheets(book.Worksheets.Count).Delete()
        for title in ("Theme", "OriginalOrange", "Manual", "Renamed"):
            book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count)).Name = title
        for sheet in book.Worksheets:
            sheet.Range("A1").Value2 = sheet.Name
        book.Worksheets("Theme").Tab.ThemeColor = 5
        book.Worksheets("Theme").Tab.TintAndShade = 0.4
        book.Worksheets("OriginalOrange").Tab.Color = repair._excel_color("#FFC000")
        book.Worksheets("Manual").Tab.Color = repair._excel_color("#112233")
        book.Worksheets("Renamed").Tab.Color = repair._excel_color("#123456")
        book.SaveAs(str(old / name), FileFormat=file_format)
        expected = {sheet.Name:batch._read_tab(sheet) for sheet in book.Worksheets}
        for sheet in book.Worksheets:
            sheet.Tab.Color = repair._excel_color("#FF0000" if sheet.Name == "Manual" else "#FFC000")
        book.Worksheets("Renamed").Name = "新Renamed"
        cell = book.Worksheets("NoColor").Range("A1")
        cell.Interior.Color = repair._excel_color("#8ACDD1")
        cell.Font.Color = repair._excel_color("#1A5276")
        orphan = book.Worksheets.Add(After=book.Worksheets(book.Worksheets.Count))
        orphan.Name = "Unpaired"
        orphan.Range("A1").Value2 = "not in backup"
        orphan.Tab.Color = repair._excel_color("#FFC000")
        book.SaveAs(str(current / name), FileFormat=file_format)
    finally:
        book.Close(SaveChanges=False)
    return expected


def verify_tint_roundtrip(excel, root):
    path, backup = root / "tint-current.xlsx", root / "tint-backup.xlsx"
    old, current = openpyxl.Workbook(), openpyxl.Workbook()
    samples = [Color(rgb=rgb, tint=tint)
               for rgb in ("123456", "4F81BD", "AA2288", "FF0000", "000000", "FFFFFF")
               for tint in (0.4, -0.4, 0.8, -0.15, 0.499984740745262, 0.5001, 0.123456)]
    samples.extend(Color(theme=index, tint=tint) for index in range(6) for tint in (0.4, -0.15))
    for row, color in enumerate(samples, 1):
        source, target = old.active.cell(row, 1, "same"), current.active.cell(row, 1, "same")
        source.fill = PatternFill("solid", fgColor=color)
        source.font = Font(color=color)
        target.fill = PatternFill("solid", fgColor="010203")
        target.font = Font(color="040506")
    old.active.sheet_properties.tabColor = Color(rgb="123456", tint=0.4)
    current.active.sheet_properties.tabColor = "010203"
    old.save(backup)
    current.save(path)
    old.close()
    current.close()
    versions = fingerprint(path), fingerprint(backup)
    book = excel.Workbooks.Open(str(backup), UpdateLinks=0, ReadOnly=True)
    xml_book = openpyxl.load_workbook(backup)
    try:
        context = color_compare._context(xml_book)
        for row, color in enumerate(samples, 1):
            cell = book.Worksheets(1).Cells(row, 1)
            expected = color_compare._color(color, context)
            assert repair._hex_of(cell.Interior.Color) == expected, (row, color, expected, cell.Interior.Color)
            assert repair._hex_of(cell.Font.Color) == expected, (row, color, expected, cell.Font.Color)
    finally:
        xml_book.close()
        book.Close(SaveChanges=False)
    planned = color_compare.plan_file(None, path, backup)["cells"]
    assert len(planned) == len(samples) + 1 and all(item["action"] == "restore" for item in planned), planned
    restored = repair.apply_file(excel, str(path), str(backup), planned, versions=versions)
    assert all(item["status"] == "ok" for item in restored), restored
    assert color_compare.plan_file(None, path, backup)["cells"] == []
    assert fingerprint(backup) == versions[1]
    return len(samples)


def main():
    with tempfile.TemporaryDirectory(prefix="tab_color_repair_") as directory:
        root = Path(directory)
        current, old = root / "current" / "001", root / "backup" / "001"
        current.mkdir(parents=True)
        old.mkdir(parents=True)
        formats = (("tabs.xlsx", 51), ("tabs.xlsm", 52), ("tabs.xls", 56), ("tabs.xlsb", 50),
                   ("tabs.xltx", 54), ("tabs.xltm", 53))
        pythoncom.CoInitialize()
        excel = None
        expected = {}
        try:
            excel = start_background_excel()
            for name, file_format in formats:
                expected[name] = create_sample(excel, current, old, name, file_format)
        finally:
            quit_background_excel(excel)
            pythoncom.CoUninitialize()
        stamps = {path:fingerprint(path) for path in old.iterdir() if not path.name.startswith("~$")}
        templates = {path:fingerprint(path) for path in current.iterdir() if repair._is_template(path)}
        with patch.object(wf, "STATE", {"busy":False, "job":None}), patch.object(wf, "log_msg"):
            repair.repair(str(current), str(old.parent), [{"fill":"#C0007A", "font":"#FFFFFF"}])
            deadline = time.monotonic() + 120
            while repair.status().get("running"):
                if time.monotonic() > deadline:
                    raise TimeoutError("标签颜色恢复超过 120 秒")
                time.sleep(0.2)
            data = repair.status()["applied"]
            assert {file["name"] for file in data["files"]} == {name for name, _ in formats}, data
            assert not any(file["error"] for file in data["files"]), data
            assert not wf.STATE["busy"]
        pythoncom.CoInitialize()
        excel = None
        try:
            excel = start_background_excel()
            tint_samples = verify_tint_roundtrip(excel, root)
            for name, _ in formats:
                if repair._is_template(name):
                    file = next(file for file in data["files"] if file["name"] == name)
                    assert not any(item.get("status") == "ok" for item in file["cells"]), file
                    assert any("模板" in item["reason"] for item in file["cells"]), file
                    continue
                book = excel.Workbooks.Open(str(current / name), UpdateLinks=0, ReadOnly=True)
                try:
                    for title in ("NoColor", "Theme", "Renamed"):
                        try:
                            target = book.Worksheets("新Renamed" if title == "Renamed" else title)
                        except Exception:
                            raise AssertionError((name, title, book.FullName, [sheet.Name for sheet in book.Worksheets])) from None
                        assert repair._tab_matches(target, expected[name][title]), (name, title, batch._read_tab(target), expected[name][title])
                    assert batch._tab_is_ours(book.Worksheets("OriginalOrange"))
                    assert batch._tab_is_ours(book.Worksheets("Unpaired"))
                    assert repair._tab_matches(book.Worksheets("Manual"), expected[name]["Manual"])
                    assert book.Worksheets("NoColor").Range("A1").Value2 == "NoColor"
                    assert book.Worksheets("NoColor").Range("A1").Interior.Pattern == -4142
                    file = next(file for file in data["files"] if file["name"] == name)
                    assert sum(item.get("kind") == "tab" and item.get("status") == "ok" for item in file["cells"]) == 4, file
                    assert any(item["sheet"] == "Unpaired" and item["action"] == "skip" for item in file["cells"]), file
                finally:
                    book.Close(SaveChanges=False)
        finally:
            quit_background_excel(excel)
            pythoncom.CoUninitialize()
        assert all(fingerprint(path) == stamp for path, stamp in stamps.items())
        assert all(fingerprint(path) == stamp for path, stamp in templates.items())
        print(json.dumps({"formats":[name for name, _ in formats], "tab_without_cell_markers":True, "templates_readonly_and_skipped":True,
            "no_color_and_theme_restored":True, "renamed_sheet_matched":True, "original_orange_kept":True,
            "arbitrary_tab_and_cell_colors_restored":True, "tab_and_cell_at_a1_restored":True, "unpaired_skipped":True,
            "backups_unchanged":True, "tint_samples":tint_samples, "tint_matches_excel":True, "no_repeated_color_repairs":True}))


if __name__ == "__main__":
    main()
