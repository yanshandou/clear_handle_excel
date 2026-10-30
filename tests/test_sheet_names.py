import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles.colors import Color

from core import excel_ops
from core.common import is_gray_tab_color


class XlsxSheetNameTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def test_reads_sheet_names_in_tab_order_without_openpyxl(self):
        wb = Workbook()
        wb.active.title = "表紙"
        wb.create_sheet("Main_20250101")
        wb.create_sheet("A&B <old>")
        wb.create_chartsheet("グラフ")
        wb.create_sheet("Main")
        wb["Main"].sheet_state = "hidden"
        path = self.folder / "book.xlsx"
        wb.save(path)

        names, backups = excel_ops.scan_workbook_sheets(path)
        self.assertEqual(names, ["表紙", "Main_20250101", "A&B <old>", "グラフ", "Main"])
        self.assertEqual(backups, ["Main_20250101"])

    def test_reads_tab_colors_rgb_theme_tint_and_indexed(self):
        wb = Workbook()
        wb.active.title = "灰"
        wb["灰"].sheet_properties.tabColor = "FF808080"
        wb.create_sheet("白色深色15%").sheet_properties.tabColor = Color(theme=0, tint=-0.1499984740745262)
        wb.create_sheet("黑色浅色50%").sheet_properties.tabColor = Color(theme=1, tint=0.499984740745262)
        wb.create_sheet("橙黄").sheet_properties.tabColor = "FFFFC000"
        wb.create_sheet("索引灰").sheet_properties.tabColor = Color(indexed=22)
        wb.create_sheet("无色")
        path = self.folder / "tabs.xlsx"
        wb.save(path)

        colors: dict[str, str] = {}
        names, _ = excel_ops.scan_workbook_sheets(path, tab_colors=colors)
        self.assertEqual(colors["灰"], "808080")
        self.assertEqual(colors["白色深色15%"], "D9D9D9")
        self.assertEqual(colors["黑色浅色50%"], "808080")
        self.assertEqual(colors["橙黄"], "FFC000")
        self.assertEqual(colors["索引灰"], "C0C0C0")
        self.assertNotIn("无色", colors)
        self.assertEqual([n for n in names if is_gray_tab_color(colors.get(n))], ["灰", "白色深色15%", "黑色浅色50%", "索引灰"])

    def test_reads_blank_sheets(self):
        # 用户 2026-10-30：只有 A1=無 的 JSONサンプル 这类没有内容的表默认删除
        wb = Workbook()
        wb.active.title = "概要"
        wb["概要"]["A1"] = "概要"
        wb["概要"]["B2"] = "説明"
        wb.create_sheet("JSONサンプル")["A1"] = "無"
        wb.create_sheet("空白")
        wb.create_sheet("書式だけ")["C3"].number_format = "@"
        wb.create_sheet("横線")["B2"] = "－"
        wb.create_sheet("該当なし")["A1"] = "該当なし。"
        wb.create_sheet("タイトルだけ")["A1"] = "帳票一覧"
        wb.create_sheet("無が二つ")["A1"] = "無"
        wb["無が二つ"]["A2"] = "無"
        wb.create_sheet("数式")["A1"] = '=""'
        wb.create_sheet("コメント")["A1"].comment = Comment("メモ", "XU")
        wb.create_sheet("数字")["A1"] = 0
        wb.create_chartsheet("グラフ")
        path = self.folder / "blank.xlsx"
        wb.save(path)

        blank: set[str] = set()
        excel_ops.scan_workbook_sheets(path, blank_sheets=blank)
        self.assertEqual(blank, {"JSONサンプル", "空白", "書式だけ", "横線", "該当なし"})

    def test_gray_tab_color_range(self):
        for color in ["808080", "F2F2F2", "0D0D0D", "D9D9D9", "E7E6E6", "A5A5A5"]:
            self.assertTrue(is_gray_tab_color(color), color)
        for color in [None, "", "000000", "FFFFFF", "44546A", "FFC000", "4472C4", "auto"]:
            self.assertFalse(is_gray_tab_color(color), color)

    def test_broken_file_still_raises(self):
        path = self.folder / "broken.xlsx"
        path.write_bytes(b"not a zip")
        with self.assertRaises(Exception):
            excel_ops.scan_workbook_sheets(path)


if __name__ == "__main__":
    unittest.main()
