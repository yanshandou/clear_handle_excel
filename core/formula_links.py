"""Excel 引用检查、公式结果转值与受保护的删表；本模块不保存工作簿。"""
from __future__ import annotations

import math
import re
from typing import Any

from openpyxl.formula.tokenizer import Tokenizer

from core import debug_log

XL_FORMULAS, XL_CONSTANTS, XL_ERRORS = -4123, 2, 16
REF_ERROR = -2146826265
ERROR_VALUES = {code - 2146828288 for code in range(2000, 2051)}
MAX_FALLBACK_CELLS = 2_000_000


def address(area: Any) -> str:
    return str(area.Address).replace("$", "")


def formula(cell: Any) -> str:
    value = cell.Formula
    return value if isinstance(value, str) and value.startswith("=") else ""


def tokens(expression: str) -> list:
    try:
        return Tokenizer(expression).items
    except Exception as error:
        raise ValueError(f"无法分析公式，未删除工作表：{expression[:160]}") from error


def broken_formula(expression: str) -> bool:
    if not expression or "#REF!" not in expression.upper():
        return False
    return any(
        (item.type == "OPERAND" and item.subtype != "TEXT" and "#REF!" in item.value.upper())
        for item in tokens(expression))


def _fallback_cells(sheet: Any, kind: int) -> list[Any]:
    used = sheet.UsedRange
    if int(used.CountLarge) > MAX_FALLBACK_CELLS:
        raise ValueError(f"「{sheet.Name}」无法读取公式，已停止处理（已用区域过大）")
    data = used.Formula if kind == XL_FORMULAS else used.Value2
    matrix = data if isinstance(data, tuple) else ((data,),)
    result = []
    for row, values in enumerate(matrix):
        for col, value in enumerate(values):
            if ((kind == XL_FORMULAS and isinstance(value, str) and value.startswith("=")) or
                    (kind == XL_CONSTANTS and isinstance(value, int) and value in ERROR_VALUES)):
                cell = sheet.Cells(int(used.Row) + row, int(used.Column) + col)
                if kind == XL_FORMULAS or not bool(cell.HasFormula):
                    result.append(cell)
    return result


def special_cells(sheet: Any, kind: int, errors_only: bool = False) -> list[Any]:
    try:
        area = sheet.UsedRange.SpecialCells(kind, XL_ERRORS) if errors_only else sheet.UsedRange.SpecialCells(kind)
        return [cell for part in area.Areas for cell in part.Cells]
    except Exception:
        # Excel 在没有符合条件的格子时也抛 1004；读取快照确认，不能把读取失败当成没有公式。
        cells = _fallback_cells(sheet, kind)
        return [cell for cell in cells if is_error(cell)] if errors_only else cells


def is_error(cell: Any) -> bool:
    value = cell.Value2
    if isinstance(value, int) and value in ERROR_VALUES:
        return bool(cell.Worksheet.Evaluate(f"ISERROR({cell.Address})"))
    return False


def is_ref_error(cell: Any) -> bool:
    return cell.Value2 == REF_ERROR and is_error(cell)


def link_errors(book: Any) -> list[dict]:
    result = []
    for sheet in book.Worksheets:
        seen = set()
        for cell in special_cells(sheet, XL_FORMULAS) + special_cells(sheet, XL_CONSTANTS, True):
            expr = formula(cell)
            broken = broken_formula(expr)
            if not (broken or is_ref_error(cell)):
                continue
            key = address(cell)
            if key in seen:
                continue
            seen.add(key)
            result.append({"sheet": sheet.Name, "addr": key, "formula": expr,
                           "error": "公式含失效引用" if broken else "计算结果为 #REF!"})
    return result


def error_positions(book: Any) -> set[tuple[str, str]]:
    return {(item["sheet"], item["addr"]) for item in link_errors(book)}


class Dependencies:
    def __init__(self, book: Any, deleting: list[str]):
        self.sheets = [str(sheet.Name).casefold() for sheet in book.Sheets]
        self.deleting = {name.casefold() for name in deleting}
        self.book_name = str(book.Name).casefold()
        self.names: dict[tuple[str, str], str] = {}
        self.tables = set()
        for item in book.Names:
            name = str(item.Name)
            if "!" in name:
                scope, name = name.rsplit("!", 1)
                scope = scope.strip("'").replace("''", "'").casefold()
            else:
                scope = ""
            self.names[scope, name.casefold()] = str(item.RefersTo)
        for sheet in book.Worksheets:
            if str(sheet.Name).casefold() in self.deleting:
                for table in sheet.ListObjects:
                    self.tables.add(str(table.Name).casefold())

    def depends(self, expression: str, scope: str, visited: frozenset = frozenset()) -> bool:
        for item in tokens(expression):
            # INDIRECT 的表名可以从别的单元格拼出来；删表前固定这类结果，不能只搜索表名。
            if item.type == "FUNC" and item.subtype == "OPEN" and item.value.upper().split(".")[-1] == "INDIRECT(":
                return True
            if item.type != "OPERAND" or item.subtype != "RANGE":
                continue
            value = item.value
            token_scope = scope.casefold()
            if "!" in value:
                qualifier, value = value.rsplit("!", 1)
                qualifier = qualifier.strip("'").replace("''", "'")
                if "[" in qualifier:
                    match = re.search(r"\[([^]]+)\](.*)$", qualifier)
                    if not match or match[1].casefold() != self.book_name:
                        continue  # 另一份文件中的同名 Sheet 不属于本次删除。
                    qualifier = match[2]
                token_scope = qualifier.casefold()
                if ":" in token_scope:
                    first, last = token_scope.split(":", 1)
                    if first in self.sheets and last in self.sheets:
                        low, high = sorted((self.sheets.index(first), self.sheets.index(last)))
                        if self.deleting.intersection(self.sheets[low:high + 1]):
                            return True
                elif token_scope in self.deleting:
                    return True
            key = (token_scope, value.casefold())
            if key not in self.names:
                key = ("", value.casefold())
            if key in self.names and key not in visited:
                if self.depends(self.names[key], token_scope, visited | {key}):
                    return True
            if value.split("[", 1)[0].casefold() in self.tables:
                return True
        return False


def region(cell: Any) -> tuple[Any, str]:
    if bool(cell.HasArray):
        return cell.CurrentArray, "array"
    try:
        if bool(cell.HasSpill):
            return cell.SpillParent.SpillingToRange, "spill"
    except Exception:
        pass  # 旧版 Excel 没有动态数组属性。
    return cell, "cell"


def valid_values(value: Any) -> bool:
    if isinstance(value, (tuple, list)):
        return all(valid_values(item) for item in value)
    return value is None or isinstance(value, (str, bool, int)) or isinstance(value, float) and math.isfinite(value)


def capture(cell: Any) -> dict:
    area, kind = region(cell)
    values = area.Value2
    if not valid_values(values):
        raise ValueError(f"{cell.Worksheet.Name}!{address(area)} 的结果不能转成固定值")
    # 用 Excel 的 ISERROR 判断，不能将恰好等于 CVErr 编码的合法整数误判为错误。
    def has_error(value: Any) -> bool:
        return any(has_error(item) for item in value) if isinstance(value, tuple) else isinstance(value, int) and value in ERROR_VALUES
    if has_error(values):
        for part in area.Cells:
            if is_error(part):
                raise ValueError(f"{cell.Worksheet.Name}!{address(part)} 本身有公式错误，未转值")
    return {"addr": address(area), "kind": kind, "value": values,
            "formula": formula(cell), "anchor": address(cell),
            "rows": int(area.Rows.Count), "cols": int(area.Columns.Count)}


def write_values(sheet: Any, item: dict) -> None:
    area = sheet.Range(item["addr"])
    value = item["value"]

    def literal(data: Any) -> Any:
        if isinstance(data, (tuple, list)):
            return tuple(literal(part) for part in data)
        return "'" + data if isinstance(data, str) else data

    # 清除完整数组/溢出区域的公式，保留格式、合并、批注等。
    if item["kind"] == "spill":
        sheet.Range(item["anchor"]).ClearContents()
    elif item["kind"] == "array":
        area.ClearContents()
    area.Value2 = literal(value)
    expected = tuple(tuple(row) for row in value) if isinstance(value, (tuple, list)) else value
    if area.Value2 != expected or area.HasFormula is not False:
        raise ValueError(f"{sheet.Name}!{item['addr']} 转值校验失败，未保存")


def delete_sheets_preserving_values(book: Any, deleting: list[str]) -> list[dict]:
    """先完整计划并转值，再删表；新增断链即抛错，调用方必须关闭且不保存。"""
    if not deleting:
        return []
    if len(deleting) >= int(book.Sheets.Count):
        raise ValueError("不能删除全部工作表")
    book.Application.CalculateFull()
    before = error_positions(book)
    index = Dependencies(book, deleting)
    plan, seen = [], set()
    for sheet in book.Worksheets:
        if str(sheet.Name).casefold() in index.deleting:
            continue
        for cell in special_cells(sheet, XL_FORMULAS):
            if index.depends(formula(cell), str(sheet.Name)):
                item = capture(cell)
                key = (sheet.Name, item["addr"])
                if key not in seen:
                    plan.append({**item, "sheet": sheet.Name})
                    seen.add(key)
    for item in plan:
        write_values(book.Worksheets(item["sheet"]), item)
        debug_log.event("🔗 删表前保留值", 位置=f"{item['sheet']}!{item['addr']}",
                        原公式=item["formula"], 保留值=debug_log.clip(item["value"]))
    for sheet in book.Worksheets:
        if str(sheet.Name).casefold() not in index.deleting:
            for cell in special_cells(sheet, XL_FORMULAS):
                if index.depends(formula(cell), str(sheet.Name)):
                    raise ValueError(f"{sheet.Name}!{address(cell)} 仍依赖待删表，未删除")
    for name in deleting:
        book.Sheets(name).Delete()
    book.Application.CalculateFull()
    new_errors = error_positions(book) - before
    if new_errors:
        locations = "、".join(f"{sheet}!{addr}" for sheet, addr in sorted(new_errors)[:10])
        raise ValueError(f"删表后出现新的引用错误：{locations}；本文件未保存")
    return plan
