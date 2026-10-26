import re
import sys
import threading
import zipfile
from collections.abc import Callable
from xml.etree import ElementTree
from pathlib import Path
from typing import Any

from core.common import (
    EXCLUDE_PDF_SHEETS,
    XL_SHEET_VISIBLE,
    XL_WORKSHEET,
    col_index_to_letters,
    find_backup_sheets,
    is_gray_tab_color,
    normalize_used_values,
)
from core import debug_log
from core.quiet_dialogs import hide_progress_dialogs, release_excel

try:
    import pythoncom
    import win32com.client as win32
except ImportError:
    print("请先安装 pywin32: pip install pywin32")
    sys.exit(1)

try:
    from openpyxl import load_workbook
except ImportError:
    print("请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)


MSO_AUTOMATION_SECURITY_FORCE_DISABLE = 3
# 跳转目标：单元格（A1）或区域（A10:H12，分页切开的合并区域）
CELL_OR_RANGE = r"[A-Za-z]+[1-9]\d*(?::[A-Za-z]+[1-9]\d*)?"
XL_A1 = 1
XL_R1C1 = -4150


def open_workbook_checked(excel: Any, path: str, **options: Any) -> Any:
    """Excel 拒绝打开时也可能返回 None，不能让后续 Close 掩盖打开失败。"""
    workbook = excel.Workbooks.Open(path, **options)
    if workbook is None:
        raise RuntimeError(
            f"Excel 未能打开工作簿（Open 返回空对象）：{path}；"
            "请检查是否有同名工作簿已打开、文件被占用或需要交互确认"
        )
    return workbook


def start_background_excel() -> Any:
    """后台用的隐藏 Excel。调用方已 CoInitialize，用完调用 quit_background_excel。

    禁用宏和事件：带宏的文件在看不见的 Excel 里弹提示框时没人能点，任务会一直卡住
    （这两个设置只对这个 Excel 进程有效，不影响用户自己的 Excel）。
    同时装上「発行中」进度框钩子。"""
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    for name, value in (("AutomationSecurity", MSO_AUTOMATION_SECURITY_FORCE_DISABLE), ("EnableEvents", False)):
        try:
            setattr(excel, name, value)
        except Exception:
            pass
    hide_progress_dialogs(excel)
    debug_log.excel_started(excel)
    return excel


def quit_background_excel(excel: Any) -> None:
    """先把主窗口挪回原位、卸下钩子，再退出 Excel"""
    if excel is None:
        return
    release_excel(excel)
    excel.Quit()


def _visible_excel() -> Any:
    """用户正在用的 Excel；后台隐藏的副本 Excel 不交给用户，另开一个可见的。调用方已 CoInitialize。"""
    try:
        excel = win32.GetActiveObject("Excel.Application")
        if not excel.Visible:
            # 后台副本可能仍在结束一次导出，不等其退出，也不把它变成用户窗口。
            excel = win32.DispatchEx("Excel.Application")
    except pythoncom.com_error as err:
        if err.args[0] != -2147221021:  # MK_E_UNAVAILABLE: 没有正在运行的 Excel
            raise
        excel = win32.DispatchEx("Excel.Application")
    return excel


def _find_or_open(excel: Any, workbook_path: Path) -> Any:
    workbook = next(
        (book for book in excel.Workbooks
         if Path(book.FullName).resolve() == workbook_path),
        None,
    )
    if workbook is None:
        excel.Visible = True
        workbook = excel.Workbooks.Open(str(workbook_path), UpdateLinks=0, IgnoreReadOnlyRecommended=True)
    return workbook


def _bring_to_front(excel: Any) -> None:
    """把 Excel 窗口提到浏览器前面；系统不让抢前台时只在任务栏闪一下，不影响打开。"""
    if sys.platform != "win32":
        return
    try:
        import ctypes

        user32 = ctypes.WinDLL("user32")
        hwnd = int(excel.Hwnd)
        if user32.IsIconic(hwnd):
            user32.ShowWindow(hwnd, 9)  # SW_RESTORE
        # 先按一下 Alt，系统才允许后台进程把别的窗口切到前台
        user32.keybd_event(0x12, 0, 0, 0)
        user32.keybd_event(0x12, 0, 2, 0)  # KEYEVENTF_KEYUP
        user32.SetForegroundWindow(hwnd)
    except Exception:
        pass


XL_FORMULAS = -4123
XL_PART = 2
XL_BY_COLUMNS = 2
XL_PREVIOUS = 2
MIN_ZOOM = 10  # Excel 允许的最小缩放
MAX_FIT_ZOOM = 100  # 表格很窄时也不放大，保持平时的样子


def _range_right_column(target: Any) -> int:
    """区域最右一列；多块区域（打印区域可以是几块）取最右的那块。"""
    right = 0
    for area in target.Areas:
        right = max(right, int(area.Column) + int(area.Columns.Count) - 1)
    return right


def _sheet_right_column(sheet: Any) -> int:
    """表格要完整显示到哪一列：设了打印区域按打印区域，否则按最右一个有内容的格子。"""
    try:
        print_area = sheet.PageSetup.PrintArea
        if print_area:
            return _range_right_column(sheet.Range(print_area))
    except Exception:
        pass
    try:
        last = sheet.Cells.Find(
            What="*", After=sheet.Cells(1, 1), LookIn=XL_FORMULAS, LookAt=XL_PART,
            SearchOrder=XL_BY_COLUMNS, SearchDirection=XL_PREVIOUS,
        )
        if last is not None:
            return int(last.Column)
    except Exception:
        pass
    return 1


def _fits(window: Any, right_column: int) -> bool:
    visible = window.VisibleRange
    # 最后一列只露出一半也会有横向滚动，所以要看到它右边那一列才算放得下
    return int(visible.Column) + int(visible.Columns.Count) - 1 > right_column


def _goto_fit_width(excel: Any, workbook: Any, sheet: Any, address: str) -> None:
    """定位到问题格，并把缩放调到表格宽度刚好放进窗口：左边从 A 列开始，不用横向滚动就能看到问题格。

    只改窗口显示，不保存；文件原来是已保存状态的，改完缩放后照样算已保存。
    """
    target = sheet.Range(address)
    was_saved = workbook.Saved
    excel.Goto(target, Scroll=True)  # 先按原来的方式跳过去：就算下面调缩放失败，至少问题格在眼前
    try:
        window = excel.ActiveWindow
        right_column = max(_range_right_column(target), _sheet_right_column(sheet))
        usable = float(window.UsableWidth)
        width = float(sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, right_column)).Width)
        zoom = MAX_FIT_ZOOM if width <= 0 else int(usable / width * 100)
        zoom = max(MIN_ZOOM, min(MAX_FIT_ZOOM, zoom))
        window.Zoom = zoom
        window.ScrollColumn = 1
        # 估算有误差（行号栏、滚动条），实际看一下还放不下就再缩一点
        while zoom > MIN_ZOOM and not _fits(window, right_column):
            zoom = max(MIN_ZOOM, zoom - 5)
            window.Zoom = zoom
            window.ScrollColumn = 1
        window.ScrollRow = max(1, int(target.Row) - 3)  # 问题行上面留几行，看得出位置
    except Exception:
        pass  # 冻结窗格、受保护的表等调不了缩放时，保持上面已经跳过去的样子
    if was_saved:
        try:
            workbook.Saved = True
        except Exception:
            pass


def _show_in_excel(excel: Any, workbook_path: Path, sheet_name: str, cell_address: str, row: int) -> float:
    workbook = _find_or_open(excel, workbook_path)
    sheet = workbook.Sheets(sheet_name)
    excel.Visible = True
    excel.WindowState = -4143  # xlNormal
    if sheet.Visible != XL_SHEET_VISIBLE:
        raise ValueError("目标工作表已隐藏，请先在 Excel 中取消隐藏后再定位")
    workbook.Activate()
    sheet.Activate()
    if workbook.Windows.Count > 0:
        workbook.Windows(1).Visible = True
        workbook.Windows(1).WindowState = -4143  # xlNormal
    address = cell_address if re.fullmatch(CELL_OR_RANGE, cell_address) else f"A{row}"
    _goto_fit_width(excel, workbook, sheet, address)
    _bring_to_front(excel)
    return round(float(sheet.Rows(row).RowHeight or 15.0), 1)


def open_workbook_in_excel(workbook_path: Path, sheet_name: str = "", cell_address: str = "") -> str:
    """网页上点文件名时用：在可见 Excel（不是后台那个隐藏的）打开文件，知道工作表/单元格就跳过去。

    只打开、不保存。工作表不存在或被隐藏时照样打开文件，返回提示文字；一切正常返回空字符串。
    """
    workbook_path = workbook_path.resolve(strict=True)
    pythoncom.CoInitialize()
    try:
        excel = _visible_excel()
        workbook = _find_or_open(excel, workbook_path)
        excel.Visible = True
        excel.WindowState = -4143  # xlNormal
        workbook.Activate()
        if workbook.Windows.Count > 0:
            workbook.Windows(1).Visible = True
            workbook.Windows(1).WindowState = -4143  # xlNormal
        note = ""
        if sheet_name:
            try:
                sheet = workbook.Sheets(sheet_name)
            except Exception:
                sheet, note = None, f"找不到工作表「{sheet_name}」，可能已被删除或改名"
            if sheet is not None and sheet.Visible != XL_SHEET_VISIBLE:
                sheet, note = None, f"工作表「{sheet_name}」是隐藏的，没有跳过去"
            if sheet is not None:
                sheet.Activate()
                if re.fullmatch(CELL_OR_RANGE, cell_address or ""):
                    _goto_fit_width(excel, workbook, sheet, cell_address)
        _bring_to_front(excel)
        return note
    finally:
        pythoncom.CoUninitialize()


def open_workbook_for_manual_review(
    workbook_path: Path, sheet_name: str, cell_address: str, row: int
) -> float:
    """定位到可见 Excel，交给用户编辑；不保存、不关闭用户的工作簿。"""
    workbook_path = workbook_path.resolve(strict=True)
    pythoncom.CoInitialize()
    try:
        return _show_in_excel(_visible_excel(), workbook_path, sheet_name, cell_address, row)
    finally:
        pythoncom.CoUninitialize()


def open_workbooks_for_review(targets: list[tuple[Path, str, str, int]]) -> dict[Path, str]:
    """在同一个可见 Excel 里一次打开全部待复检文件，各自定位到第一处黄色单元格。

    返回打不开的文件及原因；不保存、不关闭用户的工作簿。
    """
    errors: dict[Path, str] = {}
    pythoncom.CoInitialize()
    try:
        excel = _visible_excel()
        for workbook_path, sheet_name, cell_address, row in targets:
            try:
                _show_in_excel(excel, workbook_path.resolve(strict=True), sheet_name, cell_address, row)
            except Exception as err:
                errors[workbook_path] = str(err)
        return errors
    finally:
        pythoncom.CoUninitialize()


USER_EXCEL_TIMEOUT = 30.0


def close_saved_workbooks(paths: list[Path], timeout: float = USER_EXCEL_TIMEOUT) -> dict[Path, str]:
    """统一复检前关掉用户 Excel 里已保存的待复检文件；还有未保存修改的不动，返回这些文件及原因。
    你自己的 Excel 开着对话框、或正在编辑某个单元格时不响应，程序会一直等：超过 timeout 秒就不等了，报错说明。"""
    result: dict[str, Any] = {}

    def work() -> None:
        try:
            result["left"] = _close_saved_workbooks(paths)
        except BaseException as err:  # 交回调用方的线程再抛
            result["error"] = err

    worker = threading.Thread(target=work, name="close-saved-workbooks", daemon=True)
    worker.start()
    worker.join(timeout)
    debug_log.event("🪟 关闭你 Excel 里已保存的文件", 文件数=len(paths), 超时=worker.is_alive(),
                    没保存=[Path(p).name for p in result.get("left", {})] or "-")
    if worker.is_alive():
        debug_log.dump_threads("你的 Excel 不响应")
        raise ValueError(f"你打开的 Excel {int(timeout)} 秒没有响应（可能开着对话框，或正在编辑某个单元格）。"
                         "请在 Excel 里关掉对话框、按 Enter 或 Esc 结束编辑，再点一次")
    if "error" in result:
        raise result["error"]
    return result["left"]


def _close_saved_workbooks(paths: list[Path]) -> dict[Path, str]:
    wanted = {path.resolve(): path for path in paths}
    left: dict[Path, str] = {}
    pythoncom.CoInitialize()
    try:
        try:
            excel = win32.GetActiveObject("Excel.Application")
        except pythoncom.com_error:
            return left  # 没有正在运行的 Excel
        for book in list(excel.Workbooks):
            try:
                path = Path(book.FullName).resolve()
            except Exception:
                continue
            if path not in wanted:
                continue
            debug_log.event("🪟 你的 Excel 里开着", 文件=path.name, 已保存=book.Saved)
            if book.Saved:
                book.Close(SaveChanges=False)
            else:
                left[wanted[path]] = "Excel 里还有没保存的修改，请保存后再复检"
        return left
    finally:
        pythoncom.CoUninitialize()


REVIEW_FILL_COLOR = 0x00FFFF  # 黄色（Excel 颜色是 BGR）
XL_SOLID = 1
XL_FILL_NONE = -4142
XL_GRADIENT_PATTERNS = {4000, 4001}  # 渐变填充没法原样记下来，不标黄


def _safe_get(obj: Any, name: str) -> Any:
    try:
        return getattr(obj, name)
    except Exception:
        return None


def mark_review_cells(workbook: Any, marks: list[dict[str, Any]]) -> None:
    """把每处待复检单元格（合并单元格按整个合并区域）涂成黄色，原来的填充记到 mark["fill"]。

    逐处更新 mark["state"]：marked 已标黄 / skipped 没标（附 mark["error"]）。不保存。
    """
    for mark in marks:
        try:
            if not re.fullmatch(r"[A-Za-z]+[1-9]\d*", mark["addr"]):
                raise ValueError("没有定位到具体单元格")
            area = workbook.Sheets(mark["sheet"]).Range(mark["addr"]).MergeArea
            interior = area.Cells(1, 1).Interior
            fill = {key: _safe_get(interior, key) for key in (
                "Pattern", "Color", "TintAndShade", "PatternColor", "PatternColorIndex", "PatternTintAndShade",
            )}
            if fill["Pattern"] in XL_GRADIENT_PATTERNS:
                raise ValueError("单元格是渐变填充，没有标黄")
            area.Interior.Pattern = XL_SOLID
            area.Interior.Color = REVIEW_FILL_COLOR
            mark.update(range=area.Address, fill=fill, state="marked", error="")
        except Exception as err:
            mark.update(state="skipped", error=str(err))


def restore_review_cells(workbook: Any, marks: list[dict[str, Any]]) -> None:
    """把标过黄色的单元格还原成原来的填充；用户自己改了颜色的不动。不保存。"""
    for mark in marks:
        if mark.get("state") != "marked":
            continue
        try:
            sheet = workbook.Sheets(mark["sheet"])
        except Exception:
            continue  # 工作表已被删掉（例如精简 Sheet），没有要还原的
        area = sheet.Range(mark["range"])
        current = area.Cells(1, 1).Interior
        if current.Pattern != XL_SOLID or int(current.Color or 0) != REVIEW_FILL_COLOR:
            continue
        fill = mark["fill"]
        interior = area.Interior
        if fill.get("Pattern") in (None, XL_FILL_NONE):
            interior.Pattern = XL_FILL_NONE
            continue
        interior.Pattern = fill["Pattern"]
        for key in ("Color", "TintAndShade"):
            if fill.get(key) is not None:
                setattr(interior, key, fill[key])
        if fill["Pattern"] != XL_SOLID:
            if fill.get("PatternColorIndex") == -4105:  # xlAutomatic
                interior.PatternColorIndex = -4105
            elif fill.get("PatternColor") is not None:
                interior.PatternColor = fill["PatternColor"]
            if fill.get("PatternTintAndShade") is not None:
                interior.PatternTintAndShade = fill["PatternTintAndShade"]


def _read_xlsx_sheet_names(workbook_path: Path) -> list[str]:
    """直接读取压缩包内 xl/workbook.xml 的工作表列表（顺序即标签顺序），不解析共享字符串和样式"""
    with zipfile.ZipFile(workbook_path) as archive:
        root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    ns = root.tag[: root.tag.index("}") + 1] if root.tag.startswith("{") else ""
    sheets = root.find(f"{ns}sheets")
    if sheets is None:
        raise ValueError("workbook.xml 中没有工作表列表")
    names = [el.get("name") for el in sheets.findall(f"{ns}sheet")]
    if not names or any(not name for name in names):
        raise ValueError("workbook.xml 中的工作表名称无效")
    return names


XL_COLOR_INDEX_NONE = -4142  # Tab.ColorIndex：没设标签颜色
_XML_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_XML_PKG_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_XML_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
# 主题色编号 → theme1.xml 里 clrScheme 的子元素（Excel 把前两对的深浅对调了）
_THEME_SLOTS = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                "hlink", "folHlink")
_SHEET_HEAD_BYTES = 4096  # sheetPr（含 tabColor）在工作表 XML 的最开头，只读开头这一段


def _theme_colors(archive: zipfile.ZipFile) -> dict[str, str]:
    try:
        root = ElementTree.fromstring(archive.read("xl/theme/theme1.xml"))
    except (KeyError, ElementTree.ParseError):
        return {}
    scheme = root.find(f".//{_XML_DRAWING_NS}clrScheme")
    colors: dict[str, str] = {}
    for slot in (scheme if scheme is not None else []):
        tag = slot.tag.replace(_XML_DRAWING_NS, "")
        for child in slot:
            value = child.get("lastClr") or child.get("val")
            if value and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                colors[tag] = value.upper()
    return colors


def _apply_tint(hex_color: str, tint: float) -> str:
    """Excel 的 tint：按 HLS 亮度变深（负数）或变浅（正数）"""
    import colorsys
    r, g, b = (int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    return "".join(f"{round(v * 255):02X}" for v in colorsys.hls_to_rgb(h, l, s))


def _tab_color_hex(element: Any, theme: dict[str, str]) -> str | None:
    """<tabColor rgb=… / theme=… tint=… / indexed=…> 换算成 RRGGBB；auto 或认不出时返回 None"""
    if element.get("rgb"):
        value = element.get("rgb")[-6:]
    elif element.get("theme") is not None:
        idx = int(element.get("theme"))
        value = theme.get(_THEME_SLOTS[idx]) if 0 <= idx < len(_THEME_SLOTS) else None
    elif element.get("indexed") is not None:
        from openpyxl.styles.colors import COLOR_INDEX
        idx = int(element.get("indexed"))
        value = COLOR_INDEX[idx][-6:] if 0 <= idx < len(COLOR_INDEX) else None
    else:
        value = None
    if not value or not re.fullmatch(r"[0-9A-Fa-f]{6}", value):
        return None
    tint = float(element.get("tint") or 0)
    return _apply_tint(value.upper(), tint) if tint else value.upper()


def _read_xlsx_tab_colors(workbook_path: Path) -> dict[str, str]:
    """各工作表标签颜色 {表名: RRGGBB}（没设颜色的不列出）。只读每个工作表 XML 的开头，不加载内容"""
    colors: dict[str, str] = {}
    with zipfile.ZipFile(workbook_path) as archive:
        book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        ns = book.tag[: book.tag.index("}") + 1] if book.tag.startswith("{") else ""
        rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.get("Id"): rel.get("Target") for rel in rels.findall(f"{_XML_PKG_REL_NS}Relationship")}
        theme = _theme_colors(archive)
        for sheet in book.iter(f"{ns}sheet"):
            target = targets.get(sheet.get(f"{_XML_REL_NS}id")) or ""
            part = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            try:
                with archive.open(part) as fh:
                    head = fh.read(_SHEET_HEAD_BYTES).decode("utf-8", "ignore")
            except KeyError:
                continue
            m = re.search(r"<(?:\w+:)?tabColor\b[^>]*/?>", head)
            if not m:
                continue
            element = ElementTree.fromstring(re.sub(r"<(?:\w+:)?tabColor", "<tabColor", m.group(0).rstrip("/>") + "/>"))
            hex_color = _tab_color_hex(element, theme)
            if hex_color:
                colors[sheet.get("name")] = hex_color
    return colors


def _excel_tab_colors(wb: Any) -> dict[str, str]:
    """Excel 打开的工作簿各表标签颜色 {表名: RRGGBB}（Tab.Color 已经是换算好的颜色，BGR 整数）"""
    colors: dict[str, str] = {}
    for sh in wb.Sheets:
        try:
            if int(sh.Tab.ColorIndex) == XL_COLOR_INDEX_NONE:
                continue
            bgr = int(sh.Tab.Color)
        except Exception:
            continue
        colors[sh.Name] = f"{bgr & 0xFF:02X}{(bgr >> 8) & 0xFF:02X}{(bgr >> 16) & 0xFF:02X}"
    return colors


def gray_tab_sheet_names(wb: Any) -> set[str]:
    """标签是灰色的工作表：不看名字，精简 Sheet 默认删除，导出 PDF 也跳过（用户 2026-10-26）"""
    return {name for name, color in _excel_tab_colors(wb).items() if is_gray_tab_color(color)}


def scan_workbook_sheets(
    workbook_path: Path, get_excel: Callable[[], Any] | None = None, tab_colors: dict[str, str] | None = None,
) -> tuple[list[str], list[str]]:
    """扫描工作簿中的所有 Sheet 名与其中的备份 Sheet 名（支持 .xlsx/.xlsm/.xls）

    .xls 需要用 Excel 打开；批量扫描时传入 get_excel 复用同一个 Excel 进程，避免每个文件都重启 Excel。
    传入 tab_colors（空字典）时顺便填入各表标签颜色 {表名: RRGGBB}，读不到颜色不影响扫描。
    """
    if workbook_path.suffix.lower() == ".xls":
        if get_excel is not None:
            wb = get_excel().Workbooks.Open(str(workbook_path.resolve()), ReadOnly=True, UpdateLinks=0)
            try:
                all_sheets = [sh.Name for sh in wb.Sheets]
                if tab_colors is not None:
                    tab_colors.update(_excel_tab_colors(wb))
            finally:
                wb.Close(SaveChanges=False)
            return all_sheets, find_backup_sheets(all_sheets)

        pythoncom.CoInitialize()
        excel = start_background_excel()
        wb = None
        try:
            wb = excel.Workbooks.Open(str(workbook_path.resolve()), ReadOnly=True, UpdateLinks=0)
            all_sheets = [sh.Name for sh in wb.Sheets]
            if tab_colors is not None:
                tab_colors.update(_excel_tab_colors(wb))
            return all_sheets, find_backup_sheets(all_sheets)
        finally:
            if wb is not None:
                wb.Close(SaveChanges=False)
            quit_background_excel(excel)
            pythoncom.CoUninitialize()

    try:
        all_sheets = _read_xlsx_sheet_names(workbook_path)
    except Exception:
        # 结构不标准时退回 openpyxl，由它给出可读的错误信息
        workbook = load_workbook(workbook_path, read_only=True)
        try:
            all_sheets = list(workbook.sheetnames)
        finally:
            workbook.close()
    if tab_colors is not None:
        try:
            tab_colors.update(_read_xlsx_tab_colors(workbook_path))
        except Exception as err:
            debug_log.event("⚠️ 读取标签颜色失败", 文件=workbook_path.name, 原因=f"{type(err).__name__}: {err}")
    return all_sheets, find_backup_sheets(all_sheets)


def _is_sheet_already_height_adjusted(sheet: Any) -> bool:
    try:
        for prop in sheet.CustomProperties:
            if prop.Name == "_clip_height_adjusted":
                return True
    except Exception:
        pass
    return False


def _mark_sheet_height_adjusted(sheet: Any) -> None:
    try:
        sheet.CustomProperties.Add("_clip_height_adjusted", "1")
    except Exception:
        pass


def adjust_row_heights_against_clipping(wb: Any, extra_points: float = 6.0) -> None:
    """
    智能防切字调整：
    不破坏单行表头或方眼紙格子，只针对【包含非空文本且原本行高 >= 22 磅的多行文本行】
    增加安全余量（默认 +6.0 磅），并通过自定义属性保证幂等性。
    """
    backup_names = set(find_backup_sheets([sheet.Name for sheet in wb.Sheets])) | gray_tab_sheet_names(wb)
    for sheet in wb.Sheets:
        if sheet.Visible != XL_SHEET_VISIBLE or sheet.Type != XL_WORKSHEET:
            continue
        if sheet.Name in EXCLUDE_PDF_SHEETS or sheet.Name in backup_names:
            continue
        if _is_sheet_already_height_adjusted(sheet):
            continue

        try:
            used_range = sheet.UsedRange
            row_count = used_range.Rows.Count
            if row_count > 1500:
                continue

            rows_values = normalize_used_values(used_range.Value)
            adjusted_any = False

            for r_idx in range(1, row_count + 1):
                if r_idx <= len(rows_values):
                    row_vals = rows_values[r_idx - 1]
                    if not any(cell is not None and str(cell).strip() != "" for cell in row_vals):
                        continue

                row = used_range.Rows(r_idx)
                current_h = row.RowHeight
                if current_h is not None and 22.0 <= current_h < 300.0:
                    row.RowHeight = round(float(current_h) + extra_points, 1)
                    adjusted_any = True

            if adjusted_any:
                _mark_sheet_height_adjusted(sheet)
        except Exception:
            pass


def normalize_workbook_view(
    excel: Any,
    wb: Any,
    on_error: Callable[[str], None] | None = None,
) -> None:
    """将所有可见工作表定位到左上角 A1、缩放设为 100%，并停留在第一个可见 Sheet"""
    if wb.Windows.Count > 0:
        wb.Windows(1).Visible = True

    first_visible_sheet = None

    for sheet in wb.Sheets:
        if sheet.Visible != XL_SHEET_VISIBLE:
            continue
        if first_visible_sheet is None:
            first_visible_sheet = sheet

        try:
            sheet.Activate()
            if sheet.Type == XL_WORKSHEET:
                excel.Goto(sheet.Range("A1"), Scroll=True)
            excel.ActiveWindow.Zoom = 100
        except Exception as err:
            if on_error is not None:
                on_error(f"  视图重置跳过 [{wb.Name} -> {sheet.Name}]: {err}")

    if first_visible_sheet is not None:
        try:
            first_visible_sheet.Activate()
        except Exception:
            pass


def extract_sheet_snapshot(sheet: Any) -> dict[str, Any]:
    """一次性提取工作表 UsedRange 的文本快照，避免每个切字项重复跨进程读取整表"""
    try:
        used = sheet.UsedRange
        start_r = int(used.Row)
        start_c = int(used.Column)
        rows_vals = normalize_used_values(used.Value)
    except Exception:
        start_r = 1
        start_c = 1
        rows_vals = []

    return {
        "sheet": sheet,
        "start_r": start_r,
        "start_c": start_c,
        "rows_vals": rows_vals,
        "row_heights": {},
    }


def snapshot_cells(sheet_snapshot: dict[str, Any]) -> list[tuple[int, int, Any]]:
    """快照里所有非空单元格: [(行, 列, 值)]"""
    start_r = sheet_snapshot.get("start_r", 1)
    start_c = sheet_snapshot.get("start_c", 1)
    return [
        (start_r + r_idx, start_c + c_idx, value)
        for r_idx, row in enumerate(sheet_snapshot.get("rows_vals", []))
        for c_idx, value in enumerate(row)
        if value is not None
    ]


class CellProbe:
    """按需读取单元格的显示文字和打印范围；只有和 PDF 对不上的格子才跨进程查询"""

    def __init__(self, sheet: Any):
        self.sheet = sheet
        self._hidden_rows: dict[int, bool] = {}
        self._hidden_cols: dict[int, bool] = {}
        self._areas: list[tuple[int, int, int, int]] | None = None
        self._areas_unknown = False
        self._heights: dict[int, float] = {}
        self._texts: dict[tuple[int, int], str | None] = {}

    def text(self, row: int, col: int) -> str | None:
        """显示文字；同一格只跨进程读一次。行高、列宽改变后显示可能不同，要换新的 CellProbe"""
        if (row, col) not in self._texts:
            try:
                self._texts[(row, col)] = str(self.sheet.Cells(row, col).Text)
            except Exception:
                self._texts[(row, col)] = None
        return self._texts[(row, col)]

    def row_height(self, row: int) -> float:
        """行高（隐藏行为 0）；同一行只跨进程读一次"""
        if row not in self._heights:
            try:
                r = self.sheet.Rows(row)
                self._heights[row] = 0.0 if r.Hidden else round(float(r.RowHeight or 0.0), 1)
            except Exception:
                self._heights[row] = 15.0
        return self._heights[row]

    def _print_areas(self) -> list[tuple[int, int, int, int]]:
        if self._areas is None:
            self._areas = []
            try:
                print_area = str(self.sheet.PageSetup.PrintArea or "")
                if print_area:
                    try:
                        ranges = self.sheet.Range(print_area)
                    except Exception:
                        # Excel 设成 R1C1 引用样式时 PrintArea 是「R1C1:R50C10」，Range() 认不出，先换成 A1 样式
                        converted = str(self.sheet.Application.ConvertFormula(
                            "=" + print_area, XL_R1C1, XL_A1)).lstrip("=")
                        ranges = self.sheet.Range(converted)
                    for area in ranges.Areas:
                        r0, c0 = int(area.Row), int(area.Column)
                        self._areas.append((r0, c0, r0 + int(area.Rows.Count) - 1, c0 + int(area.Columns.Count) - 1))
                    self._areas_unknown = not bool(self._areas)
            except Exception:
                self._areas = []
                self._areas_unknown = True
        return self._areas

    def _row_hidden(self, row: int) -> bool:
        if row not in self._hidden_rows:
            r = self.sheet.Rows(row)
            self._hidden_rows[row] = bool(r.Hidden) or float(r.RowHeight or 0) == 0
        return self._hidden_rows[row]

    def _col_hidden(self, col: int) -> bool:
        if col not in self._hidden_cols:
            c = self.sheet.Columns(col)
            self._hidden_cols[col] = bool(c.Hidden) or float(c.ColumnWidth or 0) == 0
        return self._hidden_cols[col]

    def printed(self, row: int, col: int) -> bool | None:
        """明确隐藏/区域外为 False，可打印为 True；属性读取失败为 None，不能据此自动修改。
        合并单元格的左上角所在行列被隐藏时，只要合并区域里还有可见的行和列，文字照样会打印。"""
        try:
            hidden = self._row_hidden(row) or self._col_hidden(col)
        except Exception:
            return None
        rows, cols = [row], [col]
        if hidden:
            try:
                merge = self.sheet.Cells(row, col).MergeArea
                r0, c0, height, width = int(merge.Row), int(merge.Column), int(merge.Rows.Count), int(merge.Columns.Count)
                if height * width <= 1:
                    return False
                rows = [r for r in range(r0, r0 + height) if not self._row_hidden(r)]
                cols = [c for c in range(c0, c0 + width) if not self._col_hidden(c)]
            except Exception:
                return False  # 这一格本身明确隐藏，读不到合并信息就按隐藏处理
            if not rows or not cols:
                return False
        areas = self._print_areas()
        if self._areas_unknown:
            return None
        return not areas or any(r0 <= r <= r1 and c0 <= c <= c1
                                for r0, c0, r1, c1 in areas for r in rows for c in cols)


def set_font_colors(sheet: Any, colors: dict[tuple[int, int], int]) -> dict[tuple[int, int], Any]:
    """把格子的字体改成指定颜色（Excel 的 BGR 整数），返回原来的颜色用于还原；改不了的格子跳过"""
    originals: dict[tuple[int, int], Any] = {}
    for (row, col), color in colors.items():
        try:
            font = sheet.Cells(row, col).Font
            originals[(row, col)] = font.Color
            font.Color = color
        except Exception:
            continue
    return originals


def restore_font_colors(sheet: Any, originals: dict[tuple[int, int], Any]) -> None:
    """还原 set_font_colors 改过的颜色（一格多种颜色的富文本读不到原色，工作簿不保存，不影响原文件）"""
    for (row, col), color in originals.items():
        if color is None:
            continue
        try:
            sheet.Cells(row, col).Font.Color = color
        except Exception:
            pass


def match_excel_row(
    sheet_snapshot: dict[str, Any],
    text_snippet: str,
    page_index: int,
    total_pages: int,
    rel_y: float,
) -> tuple[int, str, float]:
    """
    根据 PDF 中被切断的文字片段和跨页纵向位置，反推对应的 Excel 行号与单元格地址。
    返回: (row_index, cell_address, current_row_height)
    """
    sheet = sheet_snapshot["sheet"]
    start_r = sheet_snapshot["start_r"]
    start_c = sheet_snapshot["start_c"]
    rows_vals = sheet_snapshot["rows_vals"]
    row_heights: dict[int, float] = sheet_snapshot["row_heights"]

    candidates: list[tuple[int, int]] = []
    clean_snip = text_snippet.strip()
    compact_snip = re.sub(r"\s+", "", clean_snip)

    total_r = max(1, len(rows_vals))
    for r_idx, row_tuple in enumerate(rows_vals):
        excel_r = start_r + r_idx
        for c_idx, cell_val in enumerate(row_tuple):
            if cell_val is None:
                continue
            s_val = str(cell_val)
            compact_val = re.sub(r"\s+", "", s_val)
            if (
                (clean_snip and clean_snip in s_val)
                or (len(compact_snip) >= 2 and compact_snip in compact_val)
            ):
                excel_c = start_c + c_idx
                candidates.append((excel_r, excel_c))

    global_rel_y = min(1.0, max(0.0, (page_index + rel_y) / max(1, total_pages)))
    approx_r = start_r + int(total_r * global_rel_y)

    target_c = start_c
    if candidates:
        candidates.sort(key=lambda rc: abs(rc[0] - approx_r))
        target_r, target_c = candidates[0]
    else:
        target_r = max(1, approx_r)

    cell_addr = f"{col_index_to_letters(target_c)}{target_r}"

    if target_r not in row_heights:
        try:
            row_h = float(sheet.Rows(target_r).RowHeight or 15.0)
        except Exception:
            row_h = 15.0
        row_heights[target_r] = round(row_h, 1)

    return target_r, cell_addr, row_heights[target_r]
