"""按备份比较单元格配色和工作表标签；颜色是否属于预设不影响检查。"""
from __future__ import annotations

import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import column_index_from_string, get_column_letter

from core.excel_ops import _apply_tint, _THEME_SLOTS, open_workbook_checked


def _context(book) -> dict:
    colors = {}
    if book.loaded_theme:
        root = ET.fromstring(book.loaded_theme)
        scheme = root.find(".//{*}clrScheme")
        if scheme is not None:
            for slot in scheme:
                color = next(iter(slot), None)
                colors[slot.tag.rsplit("}", 1)[-1]] = (color.get("lastClr") or color.get("val")) if color is not None else None
    return {"theme": [colors.get(slot) for slot in _THEME_SLOTS], "indexed": book._colors or COLOR_INDEX, "styles": {}}


def _color(color, context, default="000000") -> str:
    if color is None or color.type == "auto":
        return default
    if color.type == "rgb":
        value = color.rgb[-6:]
    elif color.type == "theme":
        value = context["theme"][color.theme] if color.theme < len(context["theme"]) else None
    elif color.type == "indexed":
        if color.indexed in (64, 65):
            return "000000" if color.indexed == 64 else "FFFFFF"
        value = context["indexed"][color.indexed][-6:] if color.indexed < len(context["indexed"]) else None
    else:
        value = None
    if not value or len(value) != 6:
        return default
    return _apply_tint(value, color.tint).upper() if color.tint else value.upper()


def _style(cell, context) -> tuple:
    key = cell.style_id
    if key not in context["styles"]:
        fill = cell.fill
        if not hasattr(fill, "stop") and fill.patternType is None:
            signature = (None,)
        elif hasattr(fill, "stop"):
            signature = ("gradient", fill.type, fill.degree, fill.left, fill.right, fill.top, fill.bottom,
                         tuple((stop.position, _color(stop.color, context)) for stop in fill.stop))
        else:
            signature = (fill.patternType, _color(fill.fgColor, context),
                         _color(fill.bgColor, context, "FFFFFF") if fill.patternType != "solid" else None)
        context["styles"][key] = signature, _color(cell.font.color, context)
    return context["styles"][key]


def _display(signature) -> dict:
    fill, font = signature
    return {"fill": "none" if fill[0] is None else f"#{fill[1]}" if fill[0] == "solid" else "other", "font": f"#{font}"}


def _row_contents(sheet, text) -> dict:
    rows = {}
    for (row, col), cell in sheet._cells.items():
        if not isinstance(cell, MergedCell) and cell.value not in (None, ""):
            rows.setdefault(row, {})[col] = text(cell.value)
    return rows


def _visible_cell(sheet, cell):
    if isinstance(cell, MergedCell):
        return next(area.start_cell for area in sheet.merged_cells.ranges if cell.coordinate in area)
    return cell


def _dimension_differences(sheet, source, now_context, old_context):
    """整行、整列的默认配色也检查；它们可能被单元格样式覆盖，列出人工核对。"""
    default = _style(Cell(sheet, row=1, column=1), now_context)
    old_default = _style(Cell(source, row=1, column=1), old_context) if source else None

    def entry(kind, addr, anchor, signature, original):
        item = {"kind":kind, "sheet":sheet.title, "addr":addr, "anchor":anchor, "mode":"all", **_display(signature),
                "action":"skip", "code":"dimension_default",
                "reason":"整行或整列的默认配色与备份不同，可能包含单元格独立样式，请人工对照恢复"}
        if source:
            previous = _display(original)
            item.update(old_sheet=source.title, old_fill=previous["fill"], old_font=previous["font"])
        else:
            item.update(code="sheet_unmatched", reason="备份工作表无法唯一配对，请人工核对整行或整列配色")
        return item

    def columns(ws):
        return [(dim.min or column_index_from_string(key), dim.max or column_index_from_string(key), dim)
                for key, dim in ws.column_dimensions.items()] if ws else []

    present, original = columns(sheet), columns(source)
    bounds = sorted({point for first, last, _ in [*present, *original] for point in (first, last + 1)})
    for first, end in zip(bounds, bounds[1:]):
        now = next((dim for left, right, dim in reversed(present) if left <= first <= right), None)
        old = next((dim for left, right, dim in reversed(original) if left <= first <= right), None)
        signature = _style(now, now_context) if now is not None else default
        previous = _style(old, old_context) if old is not None else old_default
        if signature != previous and (source or signature != default):
            yield entry("column_color", f"{get_column_letter(first)}:{get_column_letter(end - 1)}", f"{get_column_letter(first)}1", signature, previous)
    for row in sorted(set(sheet.row_dimensions) | (set(source.row_dimensions) if source else set())):
        now, old = sheet.row_dimensions.get(row), source.row_dimensions.get(row) if source else None
        signature = _style(now, now_context) if now is not None else default
        previous = _style(old, old_context) if old is not None else old_default
        if signature != previous and (source or signature != default):
            yield entry("row_color", f"{row}:{row}", f"A{row}", signature, previous)


def _plan_xml(path: Path, old_path: Path, limit: int | None) -> dict:
    from core import mark_finder as marks

    current = openpyxl.load_workbook(path, data_only=False)
    old = openpyxl.load_workbook(old_path, data_only=False)
    try:
        now_context, old_context = _context(current), _context(old)
        matches = marks._match_sheets(current, old)
        result = []

        def append(item):
            if limit is not None and len(result) >= limit:
                return False
            if item["action"] == "restore" and (marks._is_template(path) or marks._is_template(old_path)):
                item.update(action="skip", code="template", reason="模板格式暂不自动写回，请先另存为 .xlsx / .xlsm 工作簿后恢复颜色")
            result.append(item)
            return True

        for sheet in current.worksheets:
            source, how = matches.get(sheet.title, (None, ""))
            color = sheet.sheet_properties.tabColor
            before = "none" if color is None else f"#{_color(color, now_context)}"
            if source is None:
                if color is not None and not append({"kind":"tab", "sheet":sheet.title, "addr":"A1", "anchor":"A1",
                        "mode":"all", "fill":before, "font":"", "action":"skip", "code":"sheet_unmatched",
                        "reason":"备份工作表无法唯一配对，保留当前标签颜色"}):
                    return {"cells":result, "truncated":True}
            else:
                base = source.sheet_properties.tabColor
                after = "none" if base is None else f"#{_color(base, old_context)}"
                if before != after and not append({"kind":"tab", "sheet":sheet.title, "old_sheet":source.title,
                        "addr":"A1", "anchor":"A1", "mode":"all", "fill":before, "font":"", "old_fill":after,
                        "old_font":"", "action":"restore", "code":"pending", "reason":"标签颜色与备份不同，恢复备份原色"}):
                    return {"cells":result, "truncated":True}
            for item in _dimension_differences(sheet, source, now_context, old_context):
                if not append(item):
                    return {"cells":result, "truncated":True}
            merged = {area.start_cell.coordinate:str(area) for area in sheet.merged_cells.ranges}
            old_merged = {area.start_cell.coordinate:str(area) for area in source.merged_cells.ranges} if source else {}
            now_rows = _row_contents(sheet, marks._text)
            old_rows = _row_contents(source, marks._text) if source else {}
            positions = set(sheet._cells) | (set(source._cells) if source else set())
            for row, col in sorted(positions):
                cell = sheet.cell(row, col)
                original = source.cell(row, col) if source else None
                if isinstance(cell, MergedCell) and (original is None or isinstance(original, MergedCell)):
                    continue
                # 合并关系变化时，内部格按实际显示的左上角颜色比较，差异只列出、不覆盖。
                signature = _style(_visible_cell(sheet, cell), now_context)
                previous = _style(_visible_cell(source, original), old_context) if original is not None else None
                if signature == previous or (source is None and signature == ((None,), "000000")):
                    continue
                addr = merged.get(cell.coordinate, cell.coordinate)
                item = {"sheet":sheet.title, "addr":addr, "anchor":cell.coordinate, "mode":"all", **_display(signature),
                        "action":"restore", "code":"pending", "reason":"配色与备份不同，恢复备份底色与字色"}
                if source is None:
                    item.update(action="skip", code="sheet_unmatched", reason="备份工作表无法唯一配对，保留当前单元格颜色")
                else:
                    item["old_sheet"] = source.title
                    if previous is not None:
                        original_display = _display(previous)
                        item.update(old_fill=original_display["fill"], old_font=original_display["font"])
                    if isinstance(cell, MergedCell) or isinstance(original, MergedCell) or merged.get(cell.coordinate) != old_merged.get(cell.coordinate):
                        item.update(action="skip", code="merged_range_mismatch", reason="新旧合并区域不同，请人工检查颜色对应位置")
                    elif marks._text(cell.value) != marks._text(original.value):
                        item.update(action="skip", code="content_mismatch", reason="和备份同一位置的内容或公式不同，请人工检查")
                    elif cell.value in (None, "") and now_rows.get(row, {}) != old_rows.get(row, {}):
                        item.update(action="skip", code="empty_row_mismatch", reason="空格子所在行和备份内容不同，请人工检查")
                    elif signature[0][0] == "gradient" or previous[0][0] == "gradient":
                        item.update(action="skip", code="gradient", reason="渐变填充有颜色差异，请人工对照恢复")
                if not append(item):
                    return {"cells":result, "truncated":True}
        return {"cells":result, "truncated":False}
    finally:
        current.close()
        old.close()


def plan_file(excel, path: Path, old_path: Path, limit: int | None = None) -> dict:
    from core import mark_finder as marks

    if path.suffix.lower() in marks.READABLE and old_path.suffix.lower() in marks.READABLE:
        return _plan_xml(path, old_path, limit)
    # 二进制工作簿只读另存临时副本解析颜色；原文件、备份都不保存。
    with tempfile.TemporaryDirectory(prefix="compare_colors_") as directory:
        converted = []
        for index, source in enumerate((path, old_path)):
            book = open_workbook_checked(excel, str(source.resolve()), UpdateLinks=0, ReadOnly=True)
            try:
                target = Path(directory) / f"colors-{index}.xlsx"
                book.CheckCompatibility = False
                book.SaveAs(str(target), FileFormat=51)
                converted.append(target)
            finally:
                book.Close(SaveChanges=False)
        return _plan_xml(*converted, limit)
