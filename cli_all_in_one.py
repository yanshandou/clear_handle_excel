import re
import shutil
import sys
from collections.abc import Iterable, Iterator
from datetime import datetime
from pathlib import Path
from typing import TypeVar

try:
    import win32com.client as win32
    import pythoncom
except ImportError:
    print("请先安装 pywin32: pip install pywin32")
    sys.exit(1)

try:
    from openpyxl import load_workbook
except ImportError:
    print("请先安装 openpyxl: pip install openpyxl")
    sys.exit(1)

_T = TypeVar("_T")

try:
    from tqdm import tqdm as _tqdm
except ImportError:
    _tqdm = None


def progress_iter(iterable: Iterable[_T], desc: str = "", unit: str = "") -> Iterator[_T]:
    if _tqdm is not None:
        yield from _tqdm(iterable, desc=desc, unit=unit, dynamic_ncols=True)
        return

    items = list(iterable)
    total = len(items)
    if total == 0:
        return
    bar_width = 24
    for completed, item in enumerate(items, start=1):
        filled = int(bar_width * (completed - 1) / total)
        bar = "#" * filled + "-" * (bar_width - filled)
        sys.stderr.write(f"\r{desc}: [{bar}] {completed - 1}/{total} {unit}")
        sys.stderr.flush()
        yield item
    filled = bar_width
    bar = "#" * filled
    sys.stderr.write(f"\r{desc}: [{bar}] {total}/{total} {unit}\n")
    sys.stderr.flush()


def progress_write(message: str) -> None:
    if _tqdm is not None:
        _tqdm.write(message)
        return
    sys.stderr.write(f"\n{message}\n")
    sys.stderr.flush()


DATE_PATTERN = re.compile(r"(?<!\d)((?:19|20)\d{6})(?!\d)")
EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
EXCLUDE_PDF_SHEETS = {"変更履歴"}  # 不导出为 PDF 的 Sheet 名集合
RECORD_PATH = Path(__file__).with_name("record.md")

# Excel 常量
XL_TYPE_PDF = 0
XL_SHEET_VISIBLE = -1
XL_WORKSHEET = -4167


def is_backup_sheet(sheet_name: str) -> bool:
    for match in DATE_PATTERN.finditer(sheet_name):
        try:
            datetime.strptime(match.group(1), "%Y%m%d")
            return True
        except ValueError:
            continue
    return False


def sanitize_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\r\n\t]', "_", name).strip()


def scan_workbook_sheets(workbook_path: Path) -> tuple[list[str], list[str]]:
    if workbook_path.suffix.lower() == ".xls":
        pythoncom.CoInitialize()
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = None
        try:
            wb = excel.Workbooks.Open(str(workbook_path.resolve()), ReadOnly=True)
            all_sheets = [sh.Name for sh in wb.Sheets]
            backup_sheets = [name for name in all_sheets if is_backup_sheet(name)]
            return all_sheets, backup_sheets
        finally:
            if wb is not None:
                wb.Close(SaveChanges=False)
            excel.Quit()
            pythoncom.CoUninitialize()

    workbook = load_workbook(workbook_path, read_only=True)
    try:
        all_sheets = list(workbook.sheetnames)
        backup_sheets = [name for name in all_sheets if is_backup_sheet(name)]
        return all_sheets, backup_sheets
    finally:
        workbook.close()


def append_record(
    workbook_name: str,
    original_sheet_names: list[str],
    removed_sheet_names: list[str],
) -> None:
    separator = ""
    if RECORD_PATH.exists() and RECORD_PATH.stat().st_size > 0:
        with RECORD_PATH.open("rb") as record:
            record.seek(-1, 2)
            separator = "\n" if record.read(1) == b"\n" else "\n\n"
    removed_sheets = set(removed_sheet_names)
    with RECORD_PATH.open("a", encoding="utf-8", newline="\n") as record:
        record.write(f"{separator}## {workbook_name}\n\n")
        for sheet_name in original_sheet_names:
            display_name = f"~~{sheet_name}~~" if sheet_name in removed_sheets else sheet_name
            record.write(f"- {display_name}\n")


def _normalize_used_values(vals) -> list[tuple]:
    if vals is None:
        return []
    if not isinstance(vals, tuple):
        return [(vals,)]
    if vals and not isinstance(vals[0], tuple):
        return [vals]
    return list(vals)


def _unique_pdf_path(pdf_folder: Path, file_stem: str, safe_sheet_name: str, used_paths: set[Path]) -> Path:
    candidate = pdf_folder / f"{file_stem}_{safe_sheet_name}.pdf"
    counter = 2
    while candidate in used_paths:
        candidate = pdf_folder / f"{file_stem}_{safe_sheet_name}_{counter}.pdf"
        counter += 1
    used_paths.add(candidate)
    return candidate


def _is_sheet_already_height_adjusted(sheet) -> bool:
    try:
        for prop in sheet.CustomProperties:
            if prop.Name == "_clip_height_adjusted":
                return True
    except Exception:
        pass
    return False


def _mark_sheet_height_adjusted(sheet) -> None:
    try:
        sheet.CustomProperties.Add("_clip_height_adjusted", "1")
    except Exception:
        pass


def adjust_row_heights_against_clipping(wb, extra_points: float = 6.0) -> None:
    """
    智能防切字调整：
    不破坏单行表头或方眼紙格子，只针对【包含非空文本且原本行高 >= 22 磅的多行文本行】
    增加安全余量（默认 +6.0 磅），防止导出 PDF 时因折行差异导致最后一行字母压底框。
    """
    for sheet in wb.Sheets:
        if sheet.Visible != XL_SHEET_VISIBLE or sheet.Type != XL_WORKSHEET:
            continue
        if sheet.Name in EXCLUDE_PDF_SHEETS or is_backup_sheet(sheet.Name):
            continue
        if _is_sheet_already_height_adjusted(sheet):
            continue

        try:
            used_range = sheet.UsedRange
            row_count = used_range.Rows.Count
            if row_count > 1500:  # 超大表跳过逐行遍历以保证速度
                continue

            rows_values = _normalize_used_values(used_range.Value)
            adjusted_any = False

            for r_idx in range(1, row_count + 1):
                if r_idx <= len(rows_values):
                    row_vals = rows_values[r_idx - 1]
                    if not any(cell is not None and str(cell).strip() != "" for cell in row_vals):
                        continue

                row = used_range.Rows(r_idx)
                current_h = row.RowHeight
                # 只针对高度 >= 22 磅（说明里面至少有2行以上文字）且 < 300 磅的行加高
                if current_h is not None and 22.0 <= current_h < 300.0:
                    row.RowHeight = round(float(current_h) + extra_points, 1)
                    adjusted_any = True

            if adjusted_any:
                _mark_sheet_height_adjusted(sheet)
        except Exception:
            pass


def normalize_sheets_view(excel, wb) -> None:
    """将所有可见工作表定位到左上角 A1、缩放设为 100%，并停留在第一个可见 Sheet"""
    if wb.Windows.Count > 0:
        wb.Windows(1).Visible = True

    first_visible_sheet = None

    for sheet in wb.Sheets:
        if sheet.Visible != XL_SHEET_VISIBLE:
            continue
        if first_visible_sheet is None:
            first_visible_sheet = sheet

        try:
            sheet.Activate()
            if sheet.Type == XL_WORKSHEET:
                excel.Goto(sheet.Range("A1"), Scroll=True)
            excel.ActiveWindow.Zoom = 100
        except Exception as err:
            progress_write(f"  视图重置跳过 [{wb.Name} -> {sheet.Name}]: {err}")

    if first_visible_sheet is not None:
        try:
            first_visible_sheet.Activate()
        except Exception:
            pass


def process_workbooks(
    workbooks: list[Path],
    do_delete_backups: bool,
    do_fix_clipping: bool,
    do_normalize_view: bool,
    do_export_pdf: bool,
    backup_folder: Path,
    pdf_folder: Path,
) -> None:
    pythoncom.CoInitialize()
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    backup_created = False
    used_pdf_paths: set[Path] = set()
    if do_export_pdf:
        pdf_folder.mkdir(parents=True, exist_ok=True)

    try:
        for workbook_path in progress_iter(workbooks, desc="处理工作簿", unit="文件"):
            wb = None
            file_backed_up = False
            need_save = False

            try:
                wb = excel.Workbooks.Open(str(workbook_path.resolve()))
                original_sheets = [sh.Name for sh in wb.Sheets]
                backup_sheets = [name for name in original_sheets if is_backup_sheet(name)]

                def ensure_backup():
                    nonlocal backup_created, file_backed_up
                    if not backup_created:
                        backup_folder.mkdir(parents=True, exist_ok=True)
                        backup_created = True
                    if not file_backed_up:
                        shutil.copy2(workbook_path, backup_folder / workbook_path.name)
                        file_backed_up = True

                # 1. 删除备份 Sheet
                if do_delete_backups and backup_sheets:
                    if len(backup_sheets) >= len(original_sheets):
                        progress_write(f"跳过删除（删除后将无剩余工作表）: {workbook_path.name}")
                    else:
                        ensure_backup()
                        for sheet_name in backup_sheets:
                            wb.Sheets(sheet_name).Delete()
                        need_save = True
                        try:
                            append_record(workbook_path.name, original_sheets, backup_sheets)
                        except OSError as err:
                            progress_write(f"日志写入失败 {workbook_path.name}: {err}")
                        progress_write(f"已清理: {workbook_path.name} (删除 {len(backup_sheets)} 个备份表)")

                # 2. 智能增加多行文本的行高（防 PDF 切字）
                if do_fix_clipping:
                    ensure_backup()
                    adjust_row_heights_against_clipping(wb, extra_points=6.0)
                    need_save = True

                # 3. A1 定位 + 100% 缩放
                if do_normalize_view:
                    ensure_backup()
                    normalize_sheets_view(excel, wb)
                    need_save = True

                if need_save:
                    wb.Save()

                # 4. 导出 PDF
                if do_export_pdf:
                    exported_count = 0
                    file_stem = sanitize_filename(workbook_path.stem)
                    for sheet in wb.Sheets:
                        sheet_name = sheet.Name
                        if sheet_name in EXCLUDE_PDF_SHEETS or is_backup_sheet(sheet_name):
                            continue
                        if sheet.Visible != XL_SHEET_VISIBLE:
                            continue

                        safe_sheet_name = sanitize_filename(sheet_name)
                        pdf_path = _unique_pdf_path(pdf_folder, file_stem, safe_sheet_name, used_pdf_paths)

                        try:
                            sheet.ExportAsFixedFormat(
                                Type=XL_TYPE_PDF,
                                Filename=str(pdf_path.resolve()),
                                Quality=0,
                                IncludeDocProperties=True,
                                IgnorePrintAreas=False,
                                OpenAfterPublish=False,
                            )
                            exported_count += 1
                        except Exception as pdf_err:
                            progress_write(f"  导出 PDF 失败 [{workbook_path.name} -> {sheet_name}]: {pdf_err}")

                    progress_write(f"已导出 PDF: {workbook_path.name} (共 {exported_count} 个 Sheet)")

            except Exception as error:
                progress_write(f"处理失败 {workbook_path.name}: {error}")
            finally:
                if wb is not None:
                    wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()

    if backup_created:
        print(f"\n原始文件备份位置: {backup_folder}")
    if do_export_pdf:
        print(f"PDF 导出文件夹位置: {pdf_folder}")


def main() -> None:
    folder_text = input("Excel 文件夹路径: ").strip().strip('"').strip("'")
    if not folder_text:
        return
    folder = Path(folder_text).expanduser().resolve()
    if not folder.is_dir():
        print(f"文件夹不存在: {folder}")
        return

    workbooks = sorted(
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in EXCEL_EXTENSIONS
        and not path.name.startswith("~$")
    )
    if not workbooks:
        print("该文件夹中没有找到 Excel 文件。")
        return

    pending_delete: list[tuple[Path, list[str]]] = []
    for workbook_path in progress_iter(workbooks, desc="扫描文件", unit="文件"):
        try:
            all_sheets, backup_sheets = scan_workbook_sheets(workbook_path)
            if backup_sheets:
                if len(backup_sheets) >= len(all_sheets):
                    progress_write(f"警告: {workbook_path.name} 所有工作表均为日期名，将跳过删除。")
                else:
                    pending_delete.append((workbook_path, backup_sheets))
        except Exception as error:
            progress_write(f"无法读取 {workbook_path.name}: {error}")

    do_delete = False
    if pending_delete:
        print("\n检测到以下包含 YYYYMMDD 日期的备份工作表:")
        for workbook_path, sheet_names in pending_delete:
            print(f"\n{workbook_path.name}")
            for sheet_name in sheet_names:
                print(f"  - {sheet_name}")
        if input("\n[1/4] 是否删除这些备份工作表？(输入 y 确认): ").strip().lower() == "y":
            do_delete = True
    else:
        print("\n[1/4] 没有找到需要删除的 YYYYMMDD 备份工作表。")

    do_fix_clipping = (
        input("\n[2/4] 是否智能加高多行文本行的高度（防导出 PDF 切字）？(输入 y 确认): ")
        .strip()
        .lower()
        == "y"
    )

    do_normalize = (
        input("\n[3/4] 是否将所有工作表定位到 A1 单元格并以 100% 缩放保存？(输入 y 确认): ")
        .strip()
        .lower()
        == "y"
    )

    do_pdf = (
        input("\n[4/4] 是否将工作表导出为 PDF（排除「変更履歴」及日期备份表）？(输入 y 确认): ")
        .strip()
        .lower()
        == "y"
    )

    if not (do_delete or do_fix_clipping or do_normalize or do_pdf):
        print("已取消，未执行任何操作。")
        return

    timestamp = f"{datetime.now():%Y%m%d_%H%M%S}"
    backup_folder = folder / f"_sheet_cleanup_backup_{timestamp}"
    pdf_folder = folder / f"_exported_pdfs_{timestamp}"

    target_workbooks = (
        workbooks
        if (do_fix_clipping or do_normalize or do_pdf)
        else [item[0] for item in pending_delete]
    )

    process_workbooks(
        workbooks=target_workbooks,
        do_delete_backups=do_delete,
        do_fix_clipping=do_fix_clipping,
        do_normalize_view=do_normalize,
        do_export_pdf=do_pdf,
        backup_folder=backup_folder,
        pdf_folder=pdf_folder,
    )


if __name__ == "__main__":
    main()
